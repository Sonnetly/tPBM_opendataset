#!/usr/bin/env python3
"""Extract shared significant up GO and an explicit, provisional qPCR panel.

Run from the repository root:
    python3 ora_analysis/prioritize_qpcr_and_go.py

The original ORA catalog, concordant_51.csv and manual candidate config are
inputs. Panel order is scientific judgment, not a gene-level statistical rank.
No new enrichment test, DEG test or primer design is performed.
"""
from pathlib import Path
import argparse
import hashlib
import json

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parent
A = 'GSE245658_hippocampus_808nm'
B = 'GSE325653_prefrontal_1064nm'
FAMILIES = {
    '시냅스 조직화·구조': ['GO:0050808', 'GO:0050803', 'GO:0050807', 'GO:0099173', 'GO:0099084', 'GO:0099174'],
    '시냅스 신호·가소성': ['GO:0099537', 'GO:0007268', 'GO:0098916', 'GO:0048167', 'GO:0099536', 'GO:0051966'],
    '이온채널·수송·막전위': ['GO:0022836', 'GO:0022832', 'GO:0005216', 'GO:0046873', 'GO:0005244', 'GO:0022843', 'GO:0043270', 'GO:0099505', 'GO:0005261', 'GO:0060078', 'GO:0022898', 'GO:0034765', 'GO:0008066', 'GO:0016247'],
    '신경 형태·축삭 발달': ['GO:0048667', 'GO:0061564', 'GO:0007409', 'GO:0007411', 'GO:0097485', 'GO:0060560', 'GO:1990138', 'GO:0050770', 'GO:0048588'],
    '인지·학습·행동': ['GO:0050890', 'GO:0007611', 'GO:0030534'],
}
GROUP_BY_ID = {sid: family for family, ids in FAMILIES.items() for sid in ids}
assert sum(map(len, FAMILIES.values())) == len(GROUP_BY_ID) == 38


def genes(value):
    return set(filter(None, value.split(';')))


def reference(source, sid):
    return 'https://www.kegg.jp/pathway/' + sid if source == 'KEGG' else 'https://www.ebi.ac.uk/QuickGO/term/' + sid


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'qpcr_go')
    args = parser.parse_args()
    out = args.output_dir
    out.mkdir(parents=True, exist_ok=True)
    input_path = ROOT / 'input' / 'ORA_fullcatalog_KEGG_GO.csv'
    raw = pd.read_csv(input_path, keep_default_na=False, float_precision='round_trip')
    raw['input_csv_row'] = range(2, len(raw) + 2)
    assert len(raw) == 28092
    assert not raw.duplicated(['dataset', 'direction', 'source', 'set_id']).any()
    terms = {(r['dataset'], r['direction'], r['source'], r['set_id']): r for r in raw.to_dict('records')}
    concordant = pd.read_csv(ROOT / 'concordant_51.csv', keep_default_na=False)
    directions = concordant.set_index('gene')['direction'].to_dict()
    assert len(directions) == 51
    config = pd.read_csv(ROOT / 'qpcr_candidate_config.tsv', sep='\t', keep_default_na=False)
    assert not config.gene.duplicated().any() and not config.order.duplicated().any()
    recovered = pd.read_csv(ROOT / 'all_recovered_genes.csv', keep_default_na=False)
    observed = {
        r['gene']: {A: r['A_direction'] if r['A_direction'] in ['up', 'down'] else None,
                    B: r['B_direction'] if r['B_direction'] in ['up', 'down'] else None}
        for r in recovered.to_dict('records')
    }
    assert set(config.gene) <= set(observed)

    go_rows = []
    sig_go_ids = []
    up_go = raw[(raw.source == 'GO') & (raw.direction == 'up') & (raw.padj_BH_full_catalog < .05)]
    sets = {d: set(up_go[up_go.dataset == d].set_id) for d in [A, B]}
    for sid in sorted(sets[A] & sets[B]):
        a, b = [terms[d, 'up', 'GO', sid] for d in [A, B]]
        assert a['set_name'] == b['set_name'] and a['go_aspect'] == b['go_aspect']
        ga, gb = genes(a['hits']), genes(b['hits'])
        shared = ga & gb
        assert shared <= {g for g, dr in directions.items() if dr == 'up'}
        sig_go_ids.append(sid)
        go_rows.append({
            'GO_id': sid, 'GO_name': a['set_name'], 'GO_aspect': a['go_aspect'],
            'theme_manual_summary': GROUP_BY_ID[sid],
            'A_ORA_FDR': a['padj_BH_full_catalog'], 'B_ORA_FDR': b['padj_BH_full_catalog'],
            'larger_input_FDR_for_sorting': max(a['padj_BH_full_catalog'], b['padj_BH_full_catalog']),
            'shared_up_gene_count': len(shared), 'shared_up_genes': ';'.join(sorted(shared)),
            'shared_gene_support': 'has_shared_hit_genes' if shared else 'shared_term_without_shared_hit_genes',
            'A_k': a['k'], 'B_k': b['k'], 'A_fold_enrichment': a['fold_enrichment'], 'B_fold_enrichment': b['fold_enrichment'],
            'A_all_up_hits': a['hits'], 'B_all_up_hits': b['hits'],
            'A_input_csv_row': a['input_csv_row'], 'B_input_csv_row': b['input_csv_row'],
            'reference_url': reference('GO', sid),
        })
    go = pd.DataFrame(go_rows).sort_values(['larger_input_FDR_for_sorting', 'GO_id'])
    assert set(go.GO_id) == set(GROUP_BY_ID)
    assert len(go) == 38 and (go.GO_aspect == 'biological_process').sum() == 29
    assert (go.GO_aspect == 'molecular_function').sum() == 9
    assert (go.shared_up_gene_count == 0).sum() == 1
    shared_go_genes = {g for r in go_rows for g in genes(r['shared_up_genes'])}
    assert len(shared_go_genes) == 19

    baseline = pd.read_csv(ROOT / 'shared_significant_ORA.csv', keep_default_na=False, float_precision='round_trip')
    baseline_go = baseline[(baseline.source == 'GO') & (baseline.direction == 'up')].set_index('set_id')
    for r in go_rows:
        old = baseline_go.loc[r['GO_id']]
        assert old.shared_hit_genes == r['shared_up_genes']
        assert old.A_FDR == r['A_ORA_FDR'] and old.B_FDR == r['B_ORA_FDR']

    all_gene_evidence = []
    config_by_gene = config.set_index('gene').to_dict('index')
    for g in sorted(set(directions) | set(config.gene)):
        dr = directions.get(g)
        annotations = []
        for (d, direction, source, sid), a in terms.items():
            if d != A or direction != dr or g not in genes(a['hits']):
                continue
            b = terms.get((B, direction, source, sid))
            if b and g in genes(b['hits']) and a['padj_BH_full_catalog'] < .05 and b['padj_BH_full_catalog'] < .05:
                annotations.append((source, sid))
        go_ids = sorted(sid for source, sid in annotations if source == 'GO')
        kegg_ids = sorted(sid for source, sid in annotations if source == 'KEGG')
        selected = config_by_gene.get(g, {})
        all_gene_evidence.append({
            'gene': g, 'direction_from_input': dr or f"A:{observed[g][A] or 'not_observed'};B:{observed[g][B] or 'not_observed'}",
            'A_direction': observed[g][A] or 'not_observed_in_ORA_hits',
            'B_cortex_direction': observed[g][B] or 'not_observed_in_ORA_hits',
            'panel_group': selected.get('panel_group', 'not_in_proposed_panel'),
            'shared_significant_GO_count_descriptive': len(go_ids),
            'shared_significant_GO_ids': ';'.join(go_ids),
            'shared_significant_KEGG_count_descriptive': len(kegg_ids),
            'shared_significant_KEGG_ids': ';'.join(kegg_ids),
            'interpretation': 'GO count includes correlated parent/child terms; not gene-level significance or independent evidence count',
        })
    all_evidence = pd.DataFrame(all_gene_evidence)
    assert set(all_evidence.loc[(all_evidence.direction_from_input == 'up') & (all_evidence.shared_significant_GO_count_descriptive > 0), 'gene']) == shared_go_genes
    candidates = config.merge(all_evidence, on=['gene', 'panel_group'], validate='one_to_one').sort_values('order')
    assert len(candidates) == 27
    core = candidates[candidates.panel_group == 'A_cortex_priority_8']
    assert len(core) == 8
    assert set(core.gene) & shared_go_genes == {'Cacna1c', 'Grin2a', 'Atp2b3', 'Kcnj6', 'Ryr2', 'Grm2'}
    assert shared_go_genes <= set(config.gene)

    primary_rows = []
    for r in candidates.to_dict('records'):
        for sid in r['evidence_ids'].split(';'):
            source = 'GO' if sid.startswith('GO:') else 'KEGG'
            a, b = [terms.get((d, observed[r['gene']][d], source, sid)) for d in [A, B]]
            hit_a = bool(a and r['gene'] in genes(a['hits']))
            hit_b = bool(b and r['gene'] in genes(b['hits']))
            assert hit_a or hit_b, (r['gene'], sid)
            ref = a if hit_a else b
            same_direction = observed[r['gene']][A] == observed[r['gene']][B]
            primary_rows.append({
                'gene': r['gene'], 'panel_group': r['panel_group'], 'direction_from_input': r['direction_from_input'],
                'source': source, 'set_id': sid, 'set_name': ref['set_name'], 'go_aspect': ref['go_aspect'],
                'A_direction': observed[r['gene']][A] or 'not_observed_in_ORA_hits',
                'B_cortex_direction': observed[r['gene']][B] or 'not_observed_in_ORA_hits',
                'A_gene_is_hit': hit_a, 'B_gene_is_hit': hit_b,
                'A_ORA_FDR': a['padj_BH_full_catalog'] if a else None,
                'B_ORA_FDR': b['padj_BH_full_catalog'] if b else None,
                'same_term_significant_both_datasets': bool(same_direction and hit_a and hit_b and a['padj_BH_full_catalog'] < .05 and b['padj_BH_full_catalog'] < .05),
                'A_input_csv_row': a['input_csv_row'] if a else None,
                'B_input_csv_row': b['input_csv_row'] if b else None,
                'reference_url': reference(source, sid),
            })
    primary = pd.DataFrame(primary_rows)
    mitochondrial_pattern = 'mitochond|oxidative phosphorylation|respiratory chain|cellular respiration|electron transport'
    mitochondrial = raw[(raw.dataset == B) & raw.set_name.str.contains(mitochondrial_pattern, case=False, regex=True)].copy()
    oxphos = raw[(raw.dataset == B) & (raw.set_id == 'mmu00190')]
    assert len(oxphos) == 2 and (oxphos.k == 0).all() and (oxphos.padj_BH_full_catalog == 1).all()
    assert (mitochondrial.padj_BH_full_catalog < .05).sum() == 0
    core_genes = set(core.gene)
    themes = []
    for theme, ids in FAMILIES.items():
        sub = go[go.GO_id.isin(ids)]
        hit_genes = {g for s in sub.shared_up_genes for g in genes(s)}
        themes.append({
            'theme_manual_summary': theme, 'GO_term_count': len(sub),
            'shared_gene_union_count': len(hit_genes), 'shared_gene_union': ';'.join(sorted(hit_genes)),
            'core_panel_genes_in_theme': ';'.join(sorted(hit_genes & core_genes)),
            'GO_ids': ';'.join(sub.GO_id),
        })
    theme_df = pd.DataFrame(themes)
    coverage = sum(bool(genes(s) & core_genes) for s in go.shared_up_genes)
    shared_gene_sets = go[go.shared_up_gene_count > 0].groupby('shared_up_genes', sort=False).agg(
        term_count=('GO_id', 'size'), GO_ids=('GO_id', lambda s: ';'.join(s)),
        GO_names=('GO_name', lambda s: ';'.join(s)),
    ).reset_index()

    tables = {
        'qpcr_candidates': candidates, 'candidate_and_concordant_evidence': all_evidence,
        'common_up_GO_significant': go,
        'common_up_GO_BP': go[go.GO_aspect == 'biological_process'],
        'common_up_GO_MF': go[go.GO_aspect == 'molecular_function'],
        'common_up_GO_genes_19': all_evidence[all_evidence.gene.isin(shared_go_genes)].sort_values(['shared_significant_GO_count_descriptive', 'gene'], ascending=[False, True]),
        'GO_theme_summary': theme_df, 'primary_candidate_evidence': primary,
        'same_shared_hit_sets': shared_gene_sets,
        'cortex_mito_related_ORA': mitochondrial,
    }
    for name, table in tables.items():
        table.to_csv(out / (name + '.csv'), index=False, encoding='utf-8-sig')
    notes = [
        ('datasets', f'A={A}; B={B}'),
        ('selection', 'GO: up in both datasets; same GO ID; each original full-catalog ORA FDR <0.05. Gene evidence requires a shared gene in both hits lists.'),
        ('hypothesis', 'Regional differences in the sleep response to tPBM arise from differences in local cortical responsiveness to stimulation.'),
        ('panel', 'First 8 are a provisional cortical panel: 6 ion/glutamate-related common up genes + cortex-observed Ppargc1a up + shared Mrps6 down. Order is manual hypothesis-driven judgment, not a gene-level statistical ranking.'),
        ('minimal_6', 'Cacna1c;Grin2a;Atp2b3;Kcnj6;Ppargc1a;Mrps6. No guaranteed assay feasibility or validation outcome.'),
        ('mitochondrial_evidence', 'Cortical KEGG mmu00190 has k=0 and FDR=1 for both up and down. Mitochondrial GO support for Ppargc1a/Mrps6/Bcs1l is not significant. These are hypothesis-driven functional candidates, not enrichment-confirmed mitochondrial effects.'),
        ('regional_comparison', 'The two studies vary in region and wavelength, so differences cannot establish regional responsiveness. Use matched treatment/control in each cortex region and test region-by-treatment interaction; consider animal pairing and collection time/ZT.'),
        ('clock', 'Per2/Cry2 are optional clock-hypothesis targets. Cortical circadian GO is significant; Cry2 directions differ between studies. Control collection time/ZT; gene expression does not establish sleep pressure or regional specificity.'),
        ('down', 'Chrm4/Penk are optional down-direction validation targets; not reference genes and not supported by a shared significant down GO/KEGG term.'),
        ('missing', 'Gene log2FC, gene FDR, expression abundance, variance, complete DEG lists, validated qPCR assay, target sample species/tissue and contrast metadata are unavailable.'),
        ('FDR', 'A/B FDR values are original GO/KEGG term statistics. larger_input_FDR is a sort key, not a newly calculated combined FDR. No new enrichment was run.'),
        ('GO_overlap', '38 common up terms include 29 BP and 9 MF; 37 contain shared hit genes and one has zero shared genes. Union of shared hits is 19 genes.'),
        ('redundancy', 'GO annotations are correlated. Counts are descriptive. same_shared_hit_sets groups identical observed hit sets only; it is not a GO ontology or independence analysis.'),
        ('functional_specificity', 'GABA or other functional annotation may not itself be significant in both datasets. Check same_term_significant_both_datasets in primary_candidate_evidence.'),
        ('reference_genes', 'Housekeeping/reference-gene stability cannot be inferred from these ORA data or absence from hits. Validate normalization genes separately in actual samples.'),
        ('external', 'Evidence comes from uploaded ORA annotations. Official reference URLs are supplied; current databases/literature and assay specificity were not queried.'),
        ('snapshot_specific', 'Config, grouping and assertions apply to the current two datasets, not an automatic universal candidate selector.'),
        ('input_sha256', hashlib.sha256(input_path.read_bytes()).hexdigest()),
    ]
    with pd.ExcelWriter(out / 'qPCR_priority_and_common_up_GO.xlsx', engine='openpyxl') as writer:
        pd.DataFrame(notes, columns=['item', 'description']).to_excel(writer, sheet_name='READ_ME', index=False)
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name[:31], index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.fill = PatternFill('solid', fgColor='17365D')
                c.font = Font(color='FFFFFF', bold=True)
                c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[1].height = 44
            for col in ws.columns:
                header = str(col[0].value)
                ws.column_dimensions[col[0].column_letter].width = 52 if any(s in header for s in ['genes', 'name', 'reason', 'rationale', 'url', 'description', 'ids']) else 24
            for row in ws.iter_rows(min_row=2):
                for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
            url_cols = [c.column for c in ws[1] if c.value == 'reference_url']
            for idx in url_cols:
                for row in ws.iter_rows(min_row=2):
                    c = row[idx - 1]
                    c.hyperlink = c.value
                    c.style = 'Hyperlink'

    report = [
        '# qPCR 우선 후보와 공통 up 유의 GO', '',
        '> 이 보고서는 gene-level 통계를 연결하기 전의 ORA 기반 후보 제안이다. DEG FDR·log2FC를 반영한 최신 방향 검증 패널은 [DEG-aware 보고서](../deg_validation/DEG_aware_qPCR_report_ko.md)를 참고한다.', '',
        '## 사용자 가설과 우선 제안', '',
        '가설: **Regional differences in the sleep response to tPBM arise from differences in local cortical responsiveness to stimulation.** qPCR 시료는 cortex 중심이며, 영역별 전기생리 반응과 이온 신호/미토콘드리아 기전을 탐색한다.', '',
        '**1차 패널 8개: Cacna1c, Grin2a, Atp2b3, Kcnj6, Ryr2, Grm2, Ppargc1a, Mrps6.**', '',
        '앞의 6개는 두 데이터셋 공통 up인 이온/글루탐산 관련 후보다. Ppargc1a는 전전두피질 up hits에서만 관찰되고, Mrps6은 두 연구 공통 down인 미토콘드리아 번역 관련 후보다. 이 8개 전부가 공통 up이라는 뜻은 아니다.', '',
        '6개로 줄이면 Cacna1c, Grin2a, Atp2b3, Kcnj6, Ppargc1a, Mrps6을 제안한다. 12개로 확장하면 Shisa7, Slitrk3, Dlgap1, Bcs1l을 추가한다. Per2/Cry2는 일주기 가설을 별도로 검증할 때 고려할 추가 후보이고, Chrm4/Penk는 추가 down 검증 대상이다. 이들은 내부 정규화 reference gene이 아니다.', '',
        '현재 표는 ORA 결과로, log2FC/유전자별 FDR/평균 발현량/분산이 없다. 우선순위는 통계적으로 가장 강한 DEG나 실제 qPCR 검출 용이성을 뜻하지 않는다. 공통 발현 방향과 기능 대표성을 바탕으로 한 잠정적인 실험 순서다. 시료의 종·조직·처리 조건과 원본 DEG를 확인하면 조정해야 한다.', '',
        '| 순서 | 후보 | 방향 | 패널 | 선택 이유 | 대표 근거 |',
        '|---:|---|---|---|---|---|',
    ]
    for r in candidates.to_dict('records'):
        refs = []
        for sid in r['evidence_ids'].split(';'):
            e = primary[(primary.gene == r['gene']) & (primary.set_id == sid)].iloc[0]
            label = '양쪽 ORA 유의' if e.same_term_significant_both_datasets else '기능 주석; 양쪽 유의 조건 미충족'
            refs.append(f'[{sid}]({e.reference_url}) ({label})')
        report.append(f"| {r['order']} | {r['gene']} | {r['direction_from_input']} | {r['panel_group']} | {r['rationale_ko']} | {'; '.join(refs)} |")
    report += [
        '',
        '## 미토콘드리아 가설의 현재 근거', '',
        '현재 데이터에서는 이온채널/시냅스 관련 ORA 근거가 강하고 미토콘드리아 경로의 유의한 근거는 약하다. 전전두피질의 KEGG Oxidative phosphorylation(mmu00190)은 up/down 모두 k=0, FDR=1이다. 이름에 mitochond/oxidative phosphorylation/respiratory chain/cellular respiration/electron transport가 들어가는 전전두피질 ORA 항목을 추출한 부분집합에서도 FDR<0.05 항목은 없다. 이는 정의한 항목명 검색의 결과이며 모든 미토콘드리아 기능을 완전히 포괄하는 검정은 아니다.', '',
        'Ppargc1a는 전전두피질 up이며 GO:0090140(regulation of mitochondrial fission)의 FDR는 약 0.604다. Bcs1l도 전전두피질 up이며 GO:0034551(mitochondrial respiratory chain complex III assembly)의 FDR는 약 0.491이다. Mrps6은 공통 down이지만 mitochondrial translation GO의 FDR는 두 연구 모두 1이다. 이 수치는 유전자별 유의확률이 아니다. 세 후보는 데이터에서 관찰되는 가설 관련 표적이며, 유의한 미토 경로 활성/억제의 증거로 제안한 것이 아니다.', '',
        'Cox4i1/Tfam/Sod2 등 추가 기전 표적은 현재 ORA hits에서 관찰되지 않는다. 이것이 비-DEG이거나 발현하지 않음을 뜻하지는 않는다. 후보를 추가하려면 원본 DEG/발현량과 별도 주석 근거를 확인해야 한다.', '',
        '## 공통 up GO 추출 기준', '',
        '두 데이터셋의 up ORA에 같은 GO ID가 존재하고, 각 기존 full-catalog BH FDR가 모두 0.05 미만인 항목을 선택했다. 유전자 연결은 해당 GO의 두 hits 목록에 실제로 같은 유전자가 있는지로 확인했다. 최대 FDR값은 정렬 편의를 위한 값이며 새로 결합한 P/FDR가 아니다.', '',
        '**38개 항목: Biological Process 29개, Molecular Function 9개.** 이 중 공통 hit가 있는 항목은 37개, 그 유전자 합집합은 19개다. GO:1990138(neuron projection extension)은 두 데이터셋 모두 유의하지만 공통 hit는 0개다.', '',
        '공통 유의 GO의 공통 hit 19개: ' + ', '.join(sorted(shared_go_genes)) + '.', '',
        '## 기능군별 요약', '',
        '아래 묶음은 설명용 수동 분류다. GO ontology를 재구성하거나 독립적 경로로 검정한 것이 아니다.', '',
        '| 기능군 | GO 항목 수 | 공통 유전자 합집합 | 우선 패널에서 포함한 유전자 |',
        '|---|---:|---|---|',
    ]
    for r in themes:
        report.append(f"| {r['theme_manual_summary']} | {r['GO_term_count']} | {r['shared_gene_union'].replace(';', ', ')} | {r['core_panel_genes_in_theme'].replace(';', ', ')} |")
    report += ['', '## 공통 up 유의 GO 38개 전체', '',
               '| GO | 항목 | 분류 | A ORA FDR | B ORA FDR | 공통 up 유전자 |',
               '|---|---|---|---:|---:|---|']
    for r in go.to_dict('records'):
        report.append(f"| [{r['GO_id']}]({r['reference_url']}) | {r['GO_name']} | {r['GO_aspect']} | {r['A_ORA_FDR']:.4g} | {r['B_ORA_FDR']:.4g} | {r['shared_up_genes'].replace(';', ', ') or '없음: 경로만 공통'} |")
    report += [
        '', '## 실험 해석과 사전 확인', '',
        '- qPCR은 전사체 발현 방향의 재현을 확인한다. 신경 흥분성, 채널 전류, 단백질 활성, 시냅스 기능이나 PBM 인과 효과를 직접 입증하지 않는다.',
        '- 새 qPCR 시료에서 원 연구와 비교 방향·조직·조건이 대응하는지 확인한다. 다른 시료에서 같은 방향이 반드시 나올 것이라고 예측할 수는 없다.',
        '- 영역별 반응성을 보려면 같은 자극 조건에서 각 cortex 영역의 sham/control과 tPBM을 비교한다. 영역 간 baseline 차이만으로 반응성 차이를 판정하지 않는다. region×treatment 상호작용을 평가하고, 한 동물에서 여러 영역을 측정했다면 동물 내 짝지음을 분석에 반영한다.',
        '- 두 원 연구는 조직(해마/전전두피질)과 파장(808/1064 nm)이 함께 달라, 연구 간 유전자 차이로 cortex 영역별 반응성을 입증할 수 없다.',
        '- 채취 시점, 자극 후 경과 시간, 수면/각성 상태와 circadian time/ZT를 통제한다. Per2/Cry2 발현은 특히 시간 맥락과 함께 해석하며, 두 시점 또는 영역의 차이만으로 수면 반응을 단정하지 않는다.',
        '- 가능하면 원본 gene-level log2FC/FDR와 평균 발현량을 확인해 실험 표적을 조정한다. 현재 파일에서는 실제 DEG 기준을 재판정할 수 없다.',
        '- 내부 정규화 유전자는 실제 시료에서 별도로 안정성을 검증한다. GAPDH/ACTB 등의 안정성은 이 ORA 표나 hits 부재만으로 판단하지 않는다. 적절한 생물학적 반복과 assay 특이성/효율 확인이 필요하다.',
        '- 표적별 primer/assay는 아직 선정하거나 검증하지 않았다. 종과 transcript/isoform 정보를 확보한 뒤 선택해야 한다.',
        '- GO 부모·자식 항목과 동일 hit 목록은 중복될 수 있어 항목 수를 독립적 근거 수로 세지 않는다. A/B 수치는 각 연구의 기존 DEG 집합에 대한 ORA 값이며 이번 qPCR 패널만으로 새로 검정하지 않았다.',
        '- 대표 기능 주석이 있어도 그 특정 항목이 두 연구 모두에서 유의한 것은 아닐 수 있다. `primary_candidate_evidence.csv`의 유의 여부와 FDR를 함께 확인한다.',
        '- 현재 근거는 첨부 CSV의 KEGG/GO 주석이다. 공식 링크는 참조용이며 최신 외부 데이터/문헌의 직접 검증은 수행하지 못했다.',
        '', '## 재실행과 파일', '',
        '저장소 루트에서 `python3 ora_analysis/prioritize_qpcr_and_go.py`를 실행한다. 기존 ORA 분석 파일은 그대로 두고 `ora_analysis/qpcr_go/`에 출력한다.', '',
        '- `qPCR_priority_and_common_up_GO.xlsx`: 후보, 전체 GO/BP/MF, 19개 공통 유전자, 기능군, 유전자별 근거.',
        '- `qpcr_candidates.csv`: cortex 가설 기준 8개 우선 + 4개 확장 + 나머지 공통 GO 후보 10개 + Faah + 일주기 후보 2개 + 추가 down 2개.',
        '- `common_up_GO_significant.csv`: 38개 전체와 두 연구 FDR 및 실제 hits.',
        '- `common_up_GO_BP.csv`, `common_up_GO_MF.csv`: BP/MF 분리 결과.',
        '- `candidate_and_concordant_evidence.csv`: 기존 같은 방향 51개와 추가 cortex/일주기 후보의 공통 유의 항목 연결 여부 및 각 연구 관찰 방향.',
        '- `cortex_mito_related_ORA.csv`: 정의한 미토콘드리아 관련 항목명 검색으로 추출한 전전두피질 ORA 부분집합.',
        '- `primary_candidate_evidence.csv`: 수동 선정의 각 대표 근거와 원본 행 번호.',
        '- 후보 설정은 `ora_analysis/qpcr_candidate_config.tsv`에 저장했다. 다른 데이터셋에는 설정/검증 조건을 수정해야 한다.',
    ]
    (out / 'qPCR_priority_and_common_up_GO_report_ko.md').write_text('\n'.join(report) + '\n', encoding='utf-8')
    summary = {
        'input_sha256': hashlib.sha256(input_path.read_bytes()).hexdigest(),
        'shared_significant_up_GO_terms': len(go), 'BP_terms': 29, 'MF_terms': 9,
        'terms_with_shared_hit_genes': 37, 'terms_without_shared_hit_genes': 1,
        'shared_hit_genes_in_common_up_GO': len(shared_go_genes),
        'core_cortical_priority_panel': list(core.gene), 'proposed_candidate_count': len(candidates),
        'core_panel_common_up_genes': [g for g in core.gene if directions.get(g) == 'up'],
        'core_panel_other_genes': [g for g in core.gene if directions.get(g) != 'up'],
        'cortex_mito_related_ORA_name_match_count': len(mitochondrial),
        'cortex_mito_related_ORA_significant_count': int((mitochondrial.padj_BH_full_catalog < .05).sum()),
        'core_panel_shared_term_coverage_descriptive': coverage,
        'primary_evidence_rows': len(primary),
        'family_counts_descriptive': {r['theme_manual_summary']: r['GO_term_count'] for r in themes},
        'not_a_gene_level_statistical_ranking': True,
    }
    (out / 'validation_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print('OUTPUT', out)


if __name__ == '__main__':
    main()
