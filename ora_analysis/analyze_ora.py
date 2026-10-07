#!/usr/bin/env python3
"""Recover gene directions from an ORA catalog without filtering term FDR.

Usage: python3 analyze_ora.py INPUT.csv [OUTPUT_DIRECTORY]
Requires pandas, scipy, numpy, openpyxl. Does not modify the input.
"""
from pathlib import Path
import collections
import csv
import hashlib
import json
import sys

import numpy as np
import pandas as pd
from scipy.stats import false_discovery_control, hypergeom
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
INPUT = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT
OUT.mkdir(parents=True, exist_ok=True)
df = pd.read_csv(INPUT, keep_default_na=False, float_precision='round_trip')
records = df.to_dict('records')
datasets = sorted(df.dataset.unique())
assert datasets == ['GSE245658_hippocampus_808nm', 'GSE325653_prefrontal_1064nm']
A, B = datasets
assert not df.duplicated(['dataset', 'direction', 'source', 'set_id']).any()
assert set(df.direction) == {'up', 'down'}
assert set(df.source) == {'GO', 'KEGG'}
assert df.pval.between(0, 1).all() and df.padj_BH_full_catalog.between(0, 1).all()

gene_rows = collections.defaultdict(list)
group_genes = collections.defaultdict(set)
term_rows = {}
gene_annotations = collections.defaultdict(set)
for row_number, r in enumerate(records, start=2):
    r['input_csv_row'] = row_number
    r['gene_set'] = set(filter(None, r['hits'].split(';')))
    assert len(r['gene_set']) == r['k'], (row_number, 'k/hits mismatch')
    assert len(r['gene_set']) == len(list(filter(None, r['hits'].split(';'))))
    key = r['dataset'], r['direction']
    group_genes[key].update(r['gene_set'])
    term_rows[key + (r['source'], r['set_id'])] = r
    for gene in r['gene_set']:
        gene_rows[gene].append(r)
        gene_annotations[gene, *key].add((r['source'], r['set_id']))

coverage = []
numeric_checks = []
for (dataset, direction), sub in df.groupby(['dataset', 'direction']):
    assert sub.n_DEG_in_contrast.nunique() == 1
    assert sub.n_background_genes.nunique() == 1
    n = int(sub.n_DEG_in_contrast.iloc[0])
    g = group_genes[dataset, direction]
    p = hypergeom.sf(sub.k.to_numpy() - 1, sub.n_background_genes,
                     sub.set_size_tested_in_background, sub.n_DEG_in_contrast)
    q = false_discovery_control(sub.pval.to_numpy())
    assert np.allclose(p, sub.pval, rtol=1e-10, atol=1e-300)
    assert np.allclose(q, sub.padj_BH_full_catalog, rtol=1e-10, atol=1e-300)
    numeric_checks.append({
        'dataset': dataset, 'direction': direction,
        'max_abs_p_difference': float(np.max(np.abs(p - sub.pval))),
        'max_abs_BH_difference': float(np.max(np.abs(q - sub.padj_BH_full_catalog))),
    })
    coverage.append({
        'dataset': dataset, 'direction': direction, 'declared_DEG_count': n,
        'recoverable_gene_count': len(g), 'unrecoverable_gene_count': n - len(g),
        'recovered_percent': round(100 * len(g) / n, 2),
        'ORA_rows': len(sub), 'ORA_FDR_lt_005_rows': int((sub.padj_BH_full_catalog < .05).sum()),
        'background_gene_count': int(sub.n_background_genes.iloc[0]),
    })
for d in datasets:
    assert not group_genes[d, 'up'] & group_genes[d, 'down']

curated = pd.read_csv(ROOT / 'gene_function_map.tsv', sep='\t', keep_default_na=False)
curated_by_gene = curated.set_index('gene').to_dict('index')
all_genes = sorted(gene_rows)
all_gene_records = []
for g in all_genes:
    dirs = {d: next((dr for dr in ['up', 'down'] if g in group_genes[d, dr]), 'not_observed_in_ORA_hits') for d in datasets}
    both = all(dirs[d] != 'not_observed_in_ORA_hits' for d in datasets)
    status = ('same_direction' if dirs[A] == dirs[B] else 'opposite_direction') if both else ('observed_A_only' if dirs[A] != 'not_observed_in_ORA_hits' else 'observed_B_only')
    row = {'gene': g, 'A_direction': dirs[A], 'B_direction': dirs[B], 'overlap_class': status}
    for label, d in [('A', A), ('B', B)]:
        rr = [r for r in gene_rows[g] if r['dataset'] == d]
        row[label + '_GO_term_count'] = len({r['set_id'] for r in rr if r['source'] == 'GO'})
        row[label + '_KEGG_term_count'] = len({r['set_id'] for r in rr if r['source'] == 'KEGG'})
        row[label + '_in_any_FDR_lt_005_ORA_term'] = any(r['padj_BH_full_catalog'] < .05 for r in rr)
    all_gene_records.append(row)
all_gene_df = pd.DataFrame(all_gene_records)
same = all_gene_df[all_gene_df.overlap_class == 'same_direction'].copy()
opposite = all_gene_df[all_gene_df.overlap_class == 'opposite_direction'].copy()
assert set(same.gene) == set(curated.gene), (set(same.gene) ^ set(curated.gene))
assert len(same) == 51 and len(opposite) == 23
same['direction'] = same.A_direction
same = same.merge(curated, on='gene', validate='one_to_one')
same = same.sort_values(['direction', 'functional_group', 'gene'], ascending=[False, True, True])
same['KEGG_ids_in_input'] = same.gene.map(lambda g: ';'.join(sorted({r['set_id'] for r in gene_rows[g] if r['source'] == 'KEGG'})))
same['GO_ids_in_input'] = same.gene.map(lambda g: ';'.join(sorted({r['set_id'] for r in gene_rows[g] if r['source'] == 'GO'})))

def link(source, sid):
    return 'https://www.kegg.jp/pathway/' + sid if source == 'KEGG' else 'https://www.ebi.ac.uk/QuickGO/term/' + sid

evidence = []
for r in same.to_dict('records'):
    g, dr = r['gene'], r['direction']
    annotations = gene_annotations[g, A, dr] | gene_annotations[g, B, dr]
    selected = set(r['evidence_ids'].split(';'))
    assert selected <= {sid for _, sid in annotations}, (g, selected - {sid for _, sid in annotations})
    for source, sid in sorted(annotations):
        ar = term_rows.get((A, dr, source, sid))
        br = term_rows.get((B, dr, source, sid))
        ref = ar or br
        item = {'gene': g, 'direction': dr, 'functional_group': r['functional_group'],
                'source': source, 'set_id': sid, 'set_name': ref['set_name'],
                'go_aspect': ref['go_aspect'], 'selected_for_summary': sid in selected,
                'reference_url': link(source, sid),
                'evidence_origin': 'uploaded_ORA_catalog; external_reference_not_live_verified'}
        for label, tr in [('A', ar), ('B', br)]:
            hit = bool(tr and g in tr['gene_set'])
            item[label + '_gene_is_hit'] = hit
            for col in ['input_csv_row', 'k', 'n_DEG_in_contrast', 'n_background_genes', 'set_size_tested_in_background', 'fold_enrichment', 'pval', 'padj_BH_full_catalog']:
                item[label + '_' + col] = tr[col] if tr else None
        if sid in selected:
            assert item['A_gene_is_hit'] and item['B_gene_is_hit'], (g, sid, 'selected evidence not present in both datasets')
        item['ORA_FDR_lt_005_both'] = bool(ar and br and ar['padj_BH_full_catalog'] < .05 and br['padj_BH_full_catalog'] < .05)
        evidence.append(item)
evidence_df = pd.DataFrame(evidence)
primary = evidence_df[evidence_df.selected_for_summary].copy()

shared_terms = []
for dr in ['up', 'down']:
    a = {(r['source'], r['set_id']): r for r in records if r['dataset'] == A and r['direction'] == dr and r['padj_BH_full_catalog'] < .05}
    b = {(r['source'], r['set_id']): r for r in records if r['dataset'] == B and r['direction'] == dr and r['padj_BH_full_catalog'] < .05}
    for key in sorted(set(a) & set(b)):
        ar, br = a[key], b[key]
        shared_terms.append({
            'direction': dr, 'source': key[0], 'set_id': key[1], 'set_name': ar['set_name'],
            'go_aspect': ar['go_aspect'], 'A_FDR': ar['padj_BH_full_catalog'], 'B_FDR': br['padj_BH_full_catalog'],
            'A_fold_enrichment': ar['fold_enrichment'], 'B_fold_enrichment': br['fold_enrichment'],
            'shared_hit_gene_count': len(ar['gene_set'] & br['gene_set']),
            'shared_hit_genes': ';'.join(sorted(ar['gene_set'] & br['gene_set'])),
            'A_all_hits': ar['hits'], 'B_all_hits': br['hits'], 'reference_url': link(*key),
        })
shared_df = pd.DataFrame(shared_terms).sort_values(['source', 'A_FDR', 'B_FDR'])

definitions = []
for source in ['GO+KEGG', 'GO', 'KEGG']:
    for dr in ['up', 'down']:
        sig = [{g for r in records if r['dataset'] == d and r['direction'] == dr and (source == 'GO+KEGG' or r['source'] == source) and r['padj_BH_full_catalog'] < .05 for g in r['gene_set']} for d in datasets]
        definitions.append({'definition': 'genes_in_any_significant_ORA_term_in_each_dataset', 'source': source, 'direction': dr, 'count': len(sig[0] & sig[1]), 'genes': ';'.join(sorted(sig[0] & sig[1]))})
for dr in ['up', 'down']:
    genes = group_genes[A, dr] & group_genes[B, dr]
    definitions.append({'definition': 'all_ORA_hits_without_term_FDR_filter', 'source': 'GO+KEGG', 'direction': dr, 'count': len(genes), 'genes': ';'.join(sorted(genes))})
ion_genes = set(same.loc[same.functional_group == '이온채널·칼슘 수송', 'gene'])
assert len(ion_genes) == 7
channel_genes = {g for g in same.gene if any(r['set_id'] == 'GO:0005216' for r in gene_rows[g])}
assert channel_genes == ion_genes - {'Atp2b3'}
definitions.append({'definition': 'manual_ion_channel_and_calcium_transport_subset', 'source': 'functional_annotation', 'direction': 'up', 'count': len(ion_genes), 'genes': ';'.join(sorted(ion_genes))})

notes = [
    ('scope', '첨부 ORA 표 28,092행 전체 사용. ORA FDR로 유전자를 제외하지 않음. 원시 발현/count 또는 원본 DEG 표를 재분석한 결과가 아님.'),
    ('dataset_A', A), ('dataset_B', B),
    ('direction', '첨부 표의 up/down을 그대로 사용. 두 비교의 numerator/denominator, DEG 임계값, log2FC와 gene-level FDR는 제공되지 않음.'),
    ('missing_genes', 'A: 646 DEG 중 537명 복원, 109명 미복원. B: 998 DEG 중 903명 복원, 95명 미복원. missing 원인은 이 파일만으로 확정할 수 없음.'),
    ('absence', 'not_observed_in_ORA_hits는 비-DEG 또는 발현하지 않음을 뜻하지 않음. ORA hits에서 관찰되지 않았다는 뜻.'),
    ('overlap', '동일한 gene symbol의 정확한 일치로 계산. 이름이 다른 alias/ID를 추가 통합하지 않음. 같은 방향 51, 반대 방향 23. 완전한 DEG 목록 기준 최종 개수가 아님.'),
    ('FDR', 'FDR는 ORA term-level 값. 각 dataset×direction의 GO와 KEGG 전체 행을 함께 BH 보정한 값으로 수치 재현됨. DEG 유의확률이 아님.'),
    ('new_ORA', '51개 overlap 집합만을 대상으로 새 ORA를 수행하지 않음. 배경 유전자 전체 목록과 완전한 gene-set membership이 없어 적절한 재검정 불가.'),
    ('interpretation', 'up/down은 유전자 발현 방향 표기. 흥분성/억제성 신경활동의 순방향, pathway activation/inhibition, 단백질 활성 또는 인과적 PBM 효과를 뜻하지 않음.'),
    ('functional_groups', '기능군은 원본 KEGG/GO 주석을 이용한 설명용 수동 분류. 배타적인 생물학적 범주 또는 통계적 enrichment 결과가 아님.'),
    ('GO', 'GO biological_process는 생물학적 과정, molecular_function은 분자 기능. 모두를 KEGG pathway와 같은 의미로 취급하지 않음.'),
    ('references', '근거는 업로드 CSV의 hits와 set_id/set_name. 공식 KEGG/QuickGO 항목 URL 제공. 외부 KEGG/QuickGO/GEO 요청이 proxy 403으로 차단되어 최신 DB/문헌/실험조건은 미검증.'),
    ('prior_seven', '사용자는 과거 7개가 excitability/ion channel 범주였다고 설명. 현재 주석으로 채널 6개 + Atp2b3 칼슘수송을 묶으면 7개. 과거 실제 목록이 없어 동일 목록이라고 확정하지 않음.'),
    ('complete_analysis_requires', '각 데이터셋의 전체 DEG 표: gene ID/symbol, log2FC, gene-level adjusted P, comparison 정의, DEG 선정 기준. DEG 판정 재수행까지 원하면 전체 differential-expression 결과 또는 counts/design 필요.'),
    ('source_sha256', hashlib.sha256(INPUT.read_bytes()).hexdigest()),
]

outputs = {
    'coverage': pd.DataFrame(coverage), 'all_recovered_genes': all_gene_df,
    'concordant_51': same, 'discordant_23': opposite,
    'primary_evidence': primary, 'all_concordant_annotations': evidence_df,
    'shared_significant_ORA': shared_df, 'overlap_definitions': pd.DataFrame(definitions),
}
outputs['functional_groups'] = same.groupby(['direction', 'functional_group'], sort=False).agg(
    gene_count=('gene', 'size'), genes=('gene', lambda s: ';'.join(sorted(s)))
).reset_index()
for name, table in outputs.items():
    table.to_csv(OUT / (name + '.csv'), index=False, encoding='utf-8-sig')

workbook_path = OUT / 'ORA_direction_overlap_analysis.xlsx'
with pd.ExcelWriter(workbook_path, engine='openpyxl') as writer:
    pd.DataFrame(notes, columns=['item', 'description']).to_excel(writer, sheet_name='READ_ME', index=False)
    for name, table in outputs.items():
        table.to_excel(writer, sheet_name=name[:31], index=False)
    df.to_excel(writer, sheet_name='input_ORA_all_rows', index=False)
    for ws in writer.book.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[1].height = 42
        for col in ws.columns:
            header = str(col[0].value)
            width = 22
            if header in ['gene', 'direction', 'source', 'set_id']: width = 17
            if any(t in header for t in ['name', 'interpretation', 'description', 'evidence_ids', 'genes', 'hits', 'url']): width = 52
            if header == 'description': width = 115
            ws.column_dimensions[col[0].column_letter].width = width
        if ws.title in ['READ_ME', 'concordant_51', 'primary_evidence', 'discordant_23', 'shared_significant_ORA']:
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
        headers = {c.value: c.column for c in ws[1]}
        if 'reference_url' in headers:
            for row in ws.iter_rows(min_row=2):
                cell = row[headers['reference_url'] - 1]
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.style = 'Hyperlink'

stats = {
    'input_rows': len(df), 'input_sha256': hashlib.sha256(INPUT.read_bytes()).hexdigest(),
    'all_recovered_unique_genes': len(all_gene_df), 'concordant': len(same), 'discordant': len(opposite),
    'concordant_up': int((same.direction == 'up').sum()), 'concordant_down': int((same.direction == 'down').sum()),
    'direction_matrix_A_rows_B_columns': {a + '_' + b: len(group_genes[A, a] & group_genes[B, b]) for a in ['up', 'down'] for b in ['up', 'down']},
    'overlap_classes': all_gene_df.overlap_class.value_counts().to_dict(),
    'concordant_annotation_rows': len(evidence_df), 'primary_evidence_rows': len(primary),
    'shared_significant_terms': shared_df.groupby(['source', 'direction']).size().to_dict(),
    'numeric_validation': numeric_checks,
}
stats['shared_significant_terms'] = {str(k): int(v) for k, v in stats['shared_significant_terms'].items()}
(OUT / 'validation_summary.json').write_text(json.dumps(stats, indent=2, ensure_ascii=False) + '\n')

report = [
    '# ORA 전체 catalog 기반 유전자 방향 일치 분석',
    '',
    '> 이 보고서는 ORA hits에서 복원한 범위의 분석이다. 원본 DEG FDR·log2FC를 반영한 최신 overlap 및 qPCR 후보는 [최신 결과 안내](README.md)의 DEG 검증 보고서를 참고한다.',
    '',
    '**확인 가능한 같은 방향 overlap은 51개(공통 up 34개, 공통 down 17개)입니다.** 이전에 언급된 7개는 이온채널·칼슘 수송 기능군의 부분집합일 가능성이 있습니다. 기존 7개 목록 자체는 제공되지 않았습니다.',
    '',
    '이 결과는 첨부 ORA 표에서 복원되는 유전자 기준입니다. 완전한 DEG 표가 없어 전체 DEG 기준 최종 overlap을 확정할 수 없습니다. 두 데이터셋의 비교 방향은 파일의 up/down을 그대로 사용했습니다.',
    '',
    '## 사용 범위와 누락',
    '',
    '| 데이터셋 | 방향 | 표에 기록된 DEG 수 | 이름 복원 | 미복원 |',
    '|---|---|---:|---:|---:|',
]
for r in coverage:
    report.append(f"| {r['dataset']} | {r['direction']} | {r['declared_DEG_count']} | {r['recoverable_gene_count']} | {r['unrecoverable_gene_count']} |")
report += [
    '',
    '- CSV의 28,092행 전체를 사용하고, 각 hits의 세미콜론 구분 유전자 이름을 모아 중복 제거했습니다. 경로 FDR로 필터하지 않았습니다.',
    '- 복원된 유전자 중 두 데이터셋에 공통으로 있는 것은 74개: 같은 방향 51개, 반대 방향 23개입니다. 한 데이터셋에만 보이는 경우도 비-DEG라고 판정하지 않았습니다.',
    '- A에서 up/B에서 down은 12개, A에서 down/B에서 up은 11개입니다. 전체 복원 유전자의 합집합 1,366개도 엑셀에 포함했습니다.',
    '- gene-level log2FC, adjusted P, 원래 DEG 선정 기준, contrast 정의 및 전체 배경 유전자 이름은 이 CSV에 없습니다.',
    '',
    '## 과거 7개와 전체 51개의 관계',
    '',
    '**Atp2b3, Cacna1c, Grid1, Grin2a, Kcnj6, Ryr2, Trpc5**는 모두 공통 up입니다. 현재 표의 GO:0005216(ion channel activity)에 나오는 공통 유전자는 Atp2b3를 제외한 6개이며, 칼슘 수송/항상성과 KEGG calcium signaling에 연결되는 Atp2b3를 함께 묶으면 7개입니다. 이 분류가 과거의 7개와 같다는 확정은 아닙니다.',
    '',
    'FDR<0.05인 GO/KEGG 항목에 속한 유전자만 각각 모아 교집합을 구하는 정의에서는 26개(22 up, 4 down)가 나옵니다. 따라서 ‘FDR를 적용하면 항상 7개’라는 설명은 이 파일에서 재현되지 않습니다. FDR 조건과 특정 기능군 조건은 별개입니다.',
    '',
    '## 51개 유전자별 기능과 근거',
    '',
    '아래 근거는 업로드된 ORA 표의 gene–term 매핑입니다. KEGG는 경로, GO BP는 생물학적 과정, GO MF는 분자 기능입니다. 링크는 공식 항목 참조용이며 이번 환경에서는 외부 페이지를 직접 검증하지 못했습니다.',
    '',
    '| 유전자 | 공통 방향 | 기능군 | 주석에 근거한 해석 | KEGG/GO 근거 |',
    '|---|---|---|---|---|',
]
for r in same.to_dict('records'):
    links = []
    for sid in r['evidence_ids'].split(';'):
        e = primary[(primary.gene == r['gene']) & (primary.set_id == sid)].iloc[0]
        links.append(f"[{sid}]({e.reference_url}) — {e.set_name}")
    report.append(f"| {r['gene']} | {r['direction']} | {r['functional_group']} | {r['interpretation_ko']} | {'; '.join(links)} |")
report += [
    '',
    '## 경로 수준에서 상대적으로 강한 공통 근거',
    '',
    '아래 4개 KEGG 경로는 각 데이터셋의 up DEG 전체로 수행된 기존 ORA에서 모두 FDR<0.05입니다. **51개 overlap만으로 새로 검정한 enrichment가 아닙니다.** FDR는 유전자 자체의 값이 아닙니다.',
    '',
    '| KEGG 경로 | 같은 방향 공통 유전자 | A ORA FDR | B ORA FDR |',
    '|---|---|---:|---:|',
]
for r in shared_df[shared_df.source == 'KEGG'].to_dict('records'):
    report.append(f"| [{r['set_name']} ({r['set_id']})]({r['reference_url']}) | {r['shared_hit_genes'].replace(';', ', ')} | {r['A_FDR']:.3g} | {r['B_FDR']:.3g} |")
report += [
    '',
    '공통 up에서 GO 유의 항목도 38개가 겹칩니다. 예: synapse organization(GO:0050808), regulation of synaptic plasticity(GO:0048167), axon guidance(GO:0007411). GO는 부모·자식 항목이 중복되므로 38개의 독립적인 현상으로 해석하지 않습니다. 공통 유의 GO 한 항목(GO:1990138)은 두 데이터셋의 hits 교집합이 0개입니다. 경로가 공통인 것과 그 경로의 유전자가 공통인 것은 구분해야 합니다.',
    '',
    '공통 down에서는 두 데이터셋 모두 FDR<0.05인 동일 KEGG/GO 항목이 없습니다. 따라서 down 17개에 대해서는 기능적 연관성을 제시할 수 있지만, 공통적으로 유의한 경로 감소라고 표현할 수 없습니다.',
    '',
    '## 해석 시 주의할 점',
    '',
    '- 발현 up과 신경 흥분성 증가를 동일시하지 않습니다. 이온 수송 방향, 수용체 조절, 소포/시냅스 조직화 역할이 섞여 있습니다. 예를 들어 Grm2에는 adenylate cyclase-inhibiting GPCR GO가, Shisa7/Slitrk3에는 GABA 관련 GO가 함께 있습니다.',
    '- 입력 라벨상 해마 808 nm와 전전두피질 1064 nm 데이터입니다. 두 연구의 contrast, 처리 조건과 표본 구성은 확인되지 않았으므로 공통 PBM 기전이나 효과의 재현을 단정하지 않습니다.',
    '- KEGG의 심장질환·암·감염 경로에 연결되는 공유 유전자가 있어도 해당 조직에서 그 질환이 발생하거나 개선된다는 증거는 아닙니다. 본 요약은 기능적으로 직접적인 경로/GO를 우선 제시했습니다.',
    '- 51개 집합 자체의 enrichment/overlap 유의성 검정, 효과크기 비교, gene-level FDR 재판정은 수행하지 않았습니다.',
    '',
    '## 수치 검증과 산출물',
    '',
    '28,092개 행에서 k와 hits 수가 일치했고 dataset×direction별 hypergeometric P를 입력 집계값으로 재계산해 일치함을 확인했습니다. GO+KEGG 전체 catalog를 함께 BH 보정한 값도 최대 절대오차 2.3e-16 이내에서 재현되었습니다. 이는 입력된 집계값의 내부 일관성 검증이며 원래 DEG 선정이나 annotation 정확성을 외부 검증한 것은 아닙니다.',
    '',
    '- `ORA_direction_overlap_analysis.xlsx`: 전체 복원 유전자, 공통 51개, 반대방향 23개, 각 유전자에 대한 모든 주석/원본 행 번호/기존 ORA 통계, 유의 경로 비교 및 원본 전체 표.',
    '- `concordant_51.csv`: 같은 방향 51개 요약. `all_recovered_genes.csv`: 복원 가능한 모든 유전자.',
    '- `primary_evidence.csv`: 요약에 사용한 유전자별 KEGG/GO 근거와 두 데이터셋의 기존 ORA 수치.',
    '- `all_concordant_annotations.csv`: 51개에 연결된 전체 KEGG/GO 주석.',
    '- `analyze_ora.py`, `gene_function_map.tsv`: 재현용 분석 코드와 설명용 기능 분류.',
    '',
    '완전한 DEG 기준으로 마무리하려면 두 데이터셋의 원본 DEG 표(gene ID/symbol, log2FC, adjusted P와 비교 조건)가 필요합니다. 현재 결과의 미복원 유전자 이름은 ORA 표에서 추정하여 채우지 않았습니다.',
]
(OUT / 'ORA_analysis_report_ko.md').write_text('\n'.join(report) + '\n', encoding='utf-8')
print(json.dumps(stats, ensure_ascii=False, indent=2))
print('WORKBOOK', workbook_path)
