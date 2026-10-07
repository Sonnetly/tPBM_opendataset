#!/usr/bin/env python3
"""Validate the supplied FDR-selected DEG catalog and revise cortical targets.

Run from the repository root:
    python3 ora_analysis/analyze_deg_and_reprioritize.py

Statistics are imported, not recomputed from FDR-selected rows. Gene symbols
with ambiguous multiple rows retain their individual values. Different PBM
models are compared descriptively, not as a regional causal or meta-analysis.
"""
from pathlib import Path
import argparse
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parent
A, B = 'GSE245658', 'GSE325653'
A_CONTRAST = 'PBM effect: 5xFAD+light vs 5xFAD'
B_CONTRAST = 'PBM effect: hypobaric hypoxia + PBM vs hypobaric hypoxia'
PLACEHOLDERS = {'', '--', 'NA', 'N/A'}


def gene_set(value):
    return set(filter(None, value.split(';')))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'deg_validation')
    args = parser.parse_args()
    out = args.output_dir
    out.mkdir(parents=True, exist_ok=True)
    input_path = ROOT / 'input' / 'DEG_FDR05_GSE245658_GSE325653.csv'
    raw = pd.read_csv(input_path, keep_default_na=False, float_precision='round_trip',
                      dtype={'used_in_pathway_analysis': str, 'gene_symbol': str, 'gene_id': str, 'entrez': str})
    raw['input_csv_row'] = range(2, len(raw) + 2)
    assert len(raw) == 5651
    numeric = ['log2FC', 'pvalue', 'FDR_padj', 'baseMean']
    for col in numeric:
        raw[col] = pd.to_numeric(raw[col], errors='raise')
        assert np.isfinite(raw[col]).all(), col
    assert raw.FDR_padj.between(0, .05, inclusive='left').all()
    assert raw.pvalue.between(0, 1).all() and (raw.baseMean >= 0).all()
    assert (np.where(raw.log2FC > 0, 'up', 'down') == raw.direction).all()
    assert (raw.log2FC != 0).all()
    raw['pathway_eligible'] = raw.used_in_pathway_analysis.map({'True': True, 'False': False})
    assert raw.pathway_eligible.notna().all()
    raw['valid_gene_symbol'] = ~raw.gene_symbol.isin(PLACEHOLDERS)
    groups = raw.groupby(['dataset', 'tissue', 'contrast'], sort=True).agg(
        DEG_rows=('gene_symbol', 'size'), up_rows=('direction', lambda s: int((s == 'up').sum())),
        down_rows=('direction', lambda s: int((s == 'down').sum())),
        pathway_eligible_rows=('pathway_eligible', 'sum'),
        unique_valid_symbols=('gene_symbol', lambda s: len(set(s) - PLACEHOLDERS)),
    ).reset_index()
    aa = raw[(raw.dataset == A) & (raw.tissue == 'hippocampus') & (raw.contrast == A_CONTRAST)].copy()
    bb = raw[(raw.dataset == B) & (raw.tissue == 'prefrontal cortex') & (raw.contrast == B_CONTRAST)].copy()
    assert len(aa) == 657 and len(bb) == 998
    assert aa.pathway_eligible.sum() == 646 and bb.pathway_eligible.all()
    assert set(aa.wavelength_nm) == {808} and set(bb.wavelength_nm) == {1064}
    tables_by_study = {A: aa, B: bb}
    by_gene = {d: {g: sub for g, sub in frame[frame.valid_gene_symbol].groupby('gene_symbol', sort=True)} for d, frame in tables_by_study.items()}
    assert len(by_gene[A]) == 651 and len(by_gene[B]) == 998
    for d in [A, B]:
        assert all(sub.direction.nunique() == 1 for sub in by_gene[d].values()), 'Conflicting duplicate-symbol directions'
    ambiguous = pd.concat([frame[frame.duplicated('gene_symbol', keep=False)] for frame in [aa, bb]], ignore_index=True)
    placeholder = pd.concat([frame[~frame.valid_gene_symbol] for frame in [aa, bb]], ignore_index=True)

    ora = pd.read_csv(ROOT / 'input' / 'ORA_fullcatalog_KEGG_GO.csv', keep_default_na=False, float_precision='round_trip')
    ora_datasets = {A: 'GSE245658_hippocampus_808nm', B: 'GSE325653_prefrontal_1064nm'}
    audit = []
    recovered = {}
    for d in [A, B]:
        for dr in ['up', 'down']:
            sub = ora[(ora.dataset == ora_datasets[d]) & (ora.direction == dr)]
            hits = {g for value in sub.hits for g in gene_set(value)}
            recovered[d, dr] = hits
            eligible = tables_by_study[d][tables_by_study[d].pathway_eligible & (tables_by_study[d].direction == dr)]
            eligible_symbols = set(eligible.gene_symbol)
            assert eligible.gene_symbol.nunique() == len(eligible)
            assert sub.n_DEG_in_contrast.nunique() == 1
            assert int(sub.n_DEG_in_contrast.iloc[0]) == len(eligible)
            assert hits <= eligible_symbols
            audit.append({
                'dataset': d, 'direction': dr, 'ORA_declared_input_gene_count': int(sub.n_DEG_in_contrast.iloc[0]),
                'DEG_pathway_eligible_rows': len(eligible), 'ORA_hit_gene_count': len(hits),
                'eligible_DEGs_not_in_ORA_hits': len(eligible_symbols - hits),
                'not_in_ORA_hit_symbols': ';'.join(sorted(eligible_symbols - hits)),
                'input_counts_and_directions_match': True,
            })

    joined = []
    for gene in sorted(set(by_gene[A]) | set(by_gene[B])):
        dirs = {d: by_gene[d][gene].direction.iloc[0] if gene in by_gene[d] else None for d in [A, B]}
        both = dirs[A] is not None and dirs[B] is not None
        relation = ('same_direction' if dirs[A] == dirs[B] else 'opposite_direction') if both else ('A_only_in_selected_table' if dirs[A] else 'B_only_in_selected_table')
        item = {'gene': gene, 'overlap_class': relation}
        for label, d in [('A', A), ('B_cortex', B)]:
            frame = by_gene[d].get(gene)
            item[label + '_present_in_FDR05_table'] = frame is not None
            item[label + '_direction'] = dirs[d] or 'not_in_selected_FDR05_table'
            item[label + '_row_count'] = len(frame) if frame is not None else 0
            item[label + '_multiple_rows_ambiguous'] = bool(frame is not None and len(frame) > 1)
            for col in ['gene_id', 'log2FC', 'FDR_padj', 'baseMean', 'lfcSE', 'used_in_pathway_analysis', 'input_csv_row']:
                if frame is None:
                    item[label + '_' + col] = None
                elif len(frame) == 1:
                    item[label + '_' + col] = frame.iloc[0][col]
                else:
                    item[label + '_' + col] = None
                    item[label + '_' + col + '_all_values'] = ';'.join(str(x) for x in frame[col])
            item[label + '_linear_FC_treatment_over_control'] = 2 ** float(frame.log2FC.iloc[0]) if frame is not None and len(frame) == 1 else None
        item['observed_in_both_ORA_hit_lists'] = bool(both and gene in recovered[A, dirs[A]] and gene in recovered[B, dirs[B]])
        joined.append(item)
    joined_df = pd.DataFrame(joined)
    same = joined_df[joined_df.overlap_class == 'same_direction'].copy()
    opposite = joined_df[joined_df.overlap_class == 'opposite_direction'].copy()
    assert len(same) == 52 and len(opposite) == 23
    assert not same.A_multiple_rows_ambiguous.any() and not same.B_cortex_multiple_rows_ambiguous.any()
    assert (same.A_direction == 'up').sum() == 34 and (same.A_direction == 'down').sum() == 18
    old_same = pd.read_csv(ROOT / 'concordant_51.csv', keep_default_na=False)
    assert set(old_same.gene) < set(same.gene)
    assert set(same.gene) - set(old_same.gene) == {'Tmem40'}
    missing_annotation = same[~same.gene.isin(old_same.gene)].copy()

    previous = pd.read_csv(ROOT / 'qpcr_candidate_config.tsv', sep='\t', keep_default_na=False)
    priorities = pd.read_csv(ROOT / 'deg_aware_priority_config.tsv', sep='\t', keep_default_na=False)
    assert not priorities.gene.duplicated().any()
    assert set(priorities.gene) <= set(joined_df.gene)
    order = {g: i + 1 for i, g in enumerate(priorities.gene)}
    candidate_genes = set(previous.gene) | set(priorities.gene)
    candidates = joined_df[joined_df.gene.isin(candidate_genes)].copy()
    candidates = candidates.merge(previous[['gene', 'primary_focus', 'evidence_ids']], on='gene', how='left', validate='one_to_one')
    candidates = candidates.merge(priorities, on='gene', how='left', validate='one_to_one')
    candidates['priority_group'] = candidates.priority_group.fillna('E_other_previous_candidates')
    candidates['decision_reason_ko'] = candidates.decision_reason_ko.fillna('이전 기능 후보를 수치와 함께 유지; 이번 1차 패널에 선정하지 않음')
    candidates['priority_order_manual'] = candidates.gene.map(order).fillna(999).astype(int)
    candidates = candidates.sort_values(['priority_order_manual', 'B_cortex_FDR_padj', 'gene'])
    assert len(candidates) == 28 and candidates.B_cortex_present_in_FDR05_table.all()
    core = candidates[candidates.priority_group == 'A_first_8']
    assert len(core) == 8
    assert (core.B_cortex_FDR_padj.astype(float) < .05).all()
    assert (core.overlap_class == 'same_direction').all()
    stats_sorted = candidates.sort_values(['B_cortex_FDR_padj', 'gene']).copy()

    shared_go = pd.read_csv(ROOT / 'qpcr_go' / 'common_up_GO_significant.csv', keep_default_na=False, float_precision='round_trip')
    assert len(shared_go) == 38
    go_gene_pairs = []
    for r in shared_go.to_dict('records'):
        for gene in sorted(gene_set(r['shared_up_genes'])):
            original = joined_df[joined_df.gene == gene].iloc[0]
            assert original.overlap_class == 'same_direction' and original.A_direction == 'up'
            go_gene_pairs.append({
                'GO_id': r['GO_id'], 'GO_name': r['GO_name'], 'GO_aspect': r['GO_aspect'], 'gene': gene,
                'A_GO_ORA_FDR': r['A_ORA_FDR'], 'B_GO_ORA_FDR': r['B_ORA_FDR'],
                'A_gene_log2FC': original.A_log2FC, 'A_gene_FDR': original.A_FDR_padj,
                'B_cortex_gene_log2FC': original.B_cortex_log2FC, 'B_cortex_gene_FDR': original.B_cortex_FDR_padj,
            })
    go_pairs = pd.DataFrame(go_gene_pairs)
    assert go_pairs.gene.nunique() == 19

    tables = {
        'DEG_group_summary': groups, 'ORA_DEG_alignment': pd.DataFrame(audit),
        'selected_hippocampus_PBM': aa, 'selected_cortex_PBM': bb,
        'all_selected_symbol_union': joined_df, 'same_direction_52': same,
        'opposite_direction_23': opposite, 'new_unannotated_overlap': missing_annotation,
        'qPCR_candidates_with_DEG': candidates, 'candidate_cortex_FDR_sort': stats_sorted,
        'common_up_GO_with_gene_stats': go_pairs, 'common_up_GO_38': shared_go,
        'duplicate_symbols_audit': ambiguous, 'placeholder_symbols_audit': placeholder,
    }
    for name, table in tables.items():
        table.to_csv(out / (name + '.csv'), index=False, encoding='utf-8-sig')
    notes = [
        ('target_experiment', 'User plans qPCR in tPBM-treated B6 cortex. Exact B6 substrain/sex/age/time/design are not supplied. Uploaded statistics come from disease/hypoxia models, not this B6 experiment.'),
        ('scope', 'Use only hippocampus PBM effect 5xFAD+light vs 5xFAD and prefrontal PBM effect hypobaric hypoxia+PBM vs hypobaric hypoxia. Disease effects, PBM-treated vs WT and other tissues are not pooled.'),
        ('methods', 'A reports recomputed PyDESeq2 on submitter count columns; B reports submitter-provided DESeq2 statistics, not recomputed from deposited FPKM. These are file metadata, not independently verified GEO study conditions.'),
        ('FDR', 'Import supplied gene FDR from the FDR-selected file. Do not recalculate BH on the selected rows. GO/KEGG term FDR remains a separate statistic.'),
        ('overlap', 'Same direction 52 symbols: 34 up, 18 down; opposite 23. A has 657 selected rows, 651 valid unique symbols; B has 998 unique symbols. Placeholder symbols are excluded and duplicate rows retained for audit.'),
        ('ORA_input', 'A pathway-eligible rows 646 (up354/down292) and B 998 (up471/down527) match original ORA input counts exactly. Original ORA recovered 51 concordant genes; full selected DEG table adds Tmem40.'),
        ('absence', 'Not in the selected FDR05 file is not equivalent to no expression, zero effect or a known nonsignificant result. Unselected/full results are absent.'),
        ('duplicate_policy', 'Never pick the lowest FDR or largest logFC across ambiguous duplicate symbols. Individual rows and all values are retained; no duplicate affects shared 52 or the qPCR shortlist.'),
        ('priority', 'User intends direction-first candidate selection followed by qPCR verification. First 8 have matching DEG directions; FDR/logFC guide experimental order without using disease context as an exclusion rule. Cortex-only mitochondrial extras are separate. Independent per-FDR sort is included.'),
        ('baseMean', 'baseMean is supplied for assay planning; expression units/normalization are not independently verified across studies. Do not compare A/B values as a standardized expression scale. It does not predict Cq or assay feasibility.'),
        ('missing_SE', 'Cortex lfcSE is absent; no cortex confidence interval, power calculation or precision estimate is derived.'),
        ('conditions', 'Disease/hypoxia background, region, wavelength and analysis method differ. Cross-study differences cannot establish local cortical responsiveness or healthy-cortex treatment effects.'),
        ('functional_readout', 'Transcript FC is not channel activity, mitochondrial ATP production or sleep response. Test matched cortex region-by-treatment effects and connect to physiological readouts.'),
        ('source_sha256', hashlib.sha256(input_path.read_bytes()).hexdigest()),
    ]
    with pd.ExcelWriter(out / 'DEG_aware_qPCR_and_overlap.xlsx', engine='openpyxl') as writer:
        pd.DataFrame(notes, columns=['item', 'description']).to_excel(writer, sheet_name='READ_ME', index=False)
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name[:31], index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.fill = PatternFill('solid', fgColor='17365D')
                c.font = Font(color='FFFFFF', bold=True)
                c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[1].height = 45
            for col in ws.columns:
                name = str(col[0].value)
                ws.column_dimensions[col[0].column_letter].width = 60 if any(w in name for w in ['reason', 'focus', 'description', 'contrast', 'name', 'symbols']) else 25
            if ws.max_row < 300:
                for row in ws.iter_rows(min_row=2):
                    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')

    def value(x, scientific=False):
        if pd.isna(x) or x == '': return '표에 없음'
        return f'{float(x):.3g}' if scientific else f'{float(x):+.4f}'

    report = [
        '# DEG FDR·log2FC를 반영한 cortex qPCR 후보 재평가', '',
        '## 핵심 변경', '',
        '**제공된 DEG 전체에서 같은 방향 overlap은 52개(up 34/down 18)다. 이전 ORA hits 기반 51개에 Tmem40 공통 down이 추가된다. 반대 방향 overlap은 23개로 같다.**', '',
        '**B6 cortex 방향 검증 1차 8개: Grm2, Ryr2, Trpc5, Atp2b3, Cacna1c, Grin2a, Mrps6, Penk.**', '',
        '사용자는 업로드 자료의 공통 방향을 후보 선정에 활용하고, 실제 B6 cortex tPBM 시료에서 qPCR로 확인할 계획이라고 설명했다. 위 6개 공통 up은 글루탐산·이온·칼슘 축을, Mrps6/Penk 공통 down은 미토콘드리아 번역/신경펩타이드 축을 대표한다. 질환 모델 맥락을 후보 제외 조건으로 사용하지 않고, 방향 일치와 기능을 중심으로 선정했다.', '',
        'FDR와 log2FC는 qPCR 검증 순서의 참고값이다. Grin2a/Cacna1c의 기능적 중요성과 수치 근거를 함께 표시하고, Penk는 공통 down이며 변화 폭이 커서 첫 패널에 포함했다. 수동 순위는 가장 큰 fold change 순이나 FDR 순과 같지 않으며, 별도의 FDR 정렬 표도 제공한다.', '',
        '**미토콘드리아를 포함해 10개로 확장할 경우 Bcs1l/Ppargc1a를 추가한다.** 두 유전자는 피질 PBM FDR05 표에서 up이며 해마 해당 비교의 표에는 없어서 공통 방향 8개와 구분했다. Kcnj6/Slitrk3/Shisa7/Dlgap1은 추가 기전 후보로 유지한다.', '',
        'Per2/Cry2는 일주기 가설을 검증할 때 별도로 고려한다. 특히 Per2의 cortex 감소는 강한 수치 근거가 있으나, ZT/채취 시점/자극 후 시간과 수면 상태를 통제하지 않으면 해석하기 어렵다.', '',
        '## 비교 조건 확인', '',
        f'- A: {A}, hippocampus, 808 nm, **{A_CONTRAST}**.',
        f'- B: {B}, prefrontal cortex, 1064 nm, **{B_CONTRAST}**.',
        '- 전체 파일의 질병 효과, PBM-treated 5xFAD vs WT, 수막 및 림프관 내피세포 결과를 위 PBM 비교에 합치지 않았다.',
        '- 두 비교는 조직·파장뿐 아니라 5xFAD/저산소 모델 배경과 분석 방법도 다르다. 이 차이를 cortex 영역별 반응성의 직접 근거로 해석하지 않는다. 조건은 제공 파일의 metadata에서 확인한 것이며 원 GEO를 외부에서 재검증한 것은 아니다.',
        '- 목표 실험은 B6 cortex tPBM이다. B6의 정확한 substrain, 성별, 연령, 처리/채취 시간 등은 미제공이며, 원 연구의 질병 맥락을 B6 실험 조건으로 그대로 대입하지 않는다.',
        '', '## DEG와 ORA 일치 여부', '',
        '- A 해당 PBM 비교는 657행이다. 그 중 pathway-analysis 사용 표시가 True인 646행의 방향(up354/down292)이 이전 ORA 입력 수와 일치한다.',
        '- A에는 placeholder `--` 4행, Zfp809/Bpifb9a 중복 심볼이 있어, 정확한 심볼 overlap은 유효한 651개 고유 심볼로 계산했다. B는 998행/998개 고유 심볼이다. 중복의 수치 중 유리한 값을 선택하지 않았으며 audit 표에 개별 행을 보존했다.',
        '- 기존 ORA hits 537개(A)/903개(B)는 모두 올바른 조직·PBM 비교의 pathway-eligible DEG에 포함되고 방향이 일치한다.',
        '- 같은 방향 52개와 반대 방향 23개에는 ambiguous duplicate symbol이 없다. 두 표에 없는 심볼은 미발현이나 0 효과로 처리하지 않았다.',
        '', '## 1차 후보 수치', '',
        '| 후보 | 해마 log2FC | 해마 gene FDR | 피질 log2FC | 피질 gene FDR | 피질 처리/대조 fold | 선택 이유 |',
        '|---|---:|---:|---:|---:|---:|---|',
    ]
    for r in core.to_dict('records'):
        report.append(f"| {r['gene']} | {value(r['A_log2FC'])} | {value(r['A_FDR_padj'], True)} | {value(r['B_cortex_log2FC'])} | {value(r['B_cortex_FDR_padj'], True)} | {float(r['B_cortex_linear_FC_treatment_over_control']):.3f} | {r['decision_reason_ko']} |")
    report += ['', '## 확장/일주기/새 overlap 후보', '',
               '| 후보 | 분류 | 해마 log2FC | 피질 log2FC | 피질 gene FDR | 판단 |',
               '|---|---|---:|---:|---:|---|']
    for r in candidates[candidates.priority_group.isin(['B_cortex_mito_extra', 'B_mechanism_extension', 'C_clock_time_controlled', 'D_new_unannotated_overlap'])].to_dict('records'):
        report.append(f"| {r['gene']} | {r['priority_group']} | {value(r['A_log2FC'])} | {value(r['B_cortex_log2FC'])} | {value(r['B_cortex_FDR_padj'], True)} | {r['decision_reason_ko']} |")
    report += [
        '', '## GO/미토콘드리아 해석', '',
        '기존 [공통 up 유의 GO 38개](common_up_GO_38.csv)(BP29/MF9)는 이번 DEG 표로 변경되지 않는다. 해당 항목들의 실제 공통 hit 19개 모두 두 PBM 비교의 gene FDR<0.05 및 up 방향을 확인했다. [GO–gene 통계 연결 CSV](common_up_GO_with_gene_stats.csv)에는 원래 GO 항목 FDR와 각 유전자 FDR/log2FC를 별개 열로 연결했다.', '',
        '미토콘드리아 관련 Ppargc1a/Bcs1l/Mrps6은 이제 gene-level FDR<0.05라는 근거가 확인된다. 그렇더라도 미토콘드리아 GO/KEGG enrichment가 유의하다는 뜻은 아니다. 기존 cortex Oxidative phosphorylation은 hits 0, FDR1이고 대표 미토 GO 항목의 FDR도 유의하지 않다. 단일 유전자 발현으로 ATP 생산 증가/감소를 판정하지 않는다.', '',
        'Tmem40은 해마 log2FC -2.2219/FDR0.0252, 피질 -1.0335/FDR0.00112이며 공통 down이다. 이 파일들의 기존 KEGG/GO hits에는 없고 제공된 기능 주석이 없으므로 기전 중심 1차 패널에 자동 추가하지 않고 탐색 대상으로 표시한다.', '',
        '## 실험 및 수치 해석', '',
        '- 피질 1차 후보 중 많은 증가 효과는 처리/대조 약 1.2–1.4배다. 통계적으로 유의한 것과 qPCR에서 큰 변화가 예상되는 것은 다르다. assay 효율·정규화·생물학적 변동을 확인하며, Cortex lfcSE가 없어 이 표만으로 필요한 표본 수나 검출력을 계산하지 않았다.',
        '- 독립적인 cortex 영역별 control/tPBM 비교와 region×treatment 차이를 평가한다. 동일 동물의 여러 영역을 측정하면 짝지음/동물 효과를 반영한다. 채취 시점, ZT, 자극 후 시간과 수면/각성 상태를 통제한다.',
        '- qPCR은 발현을 검증하며 전기생리 반응, 채널 전류, 단백질 활성이나 미토콘드리아 호흡을 직접 측정하지 않는다. 해당 기능 readout과 연결해서 가설을 평가한다.',
        '- gene FDR는 제공 값이다. FDR<0.05로 이미 선택된 파일만으로 BH를 재계산하거나 전체 발현 배경으로 새로운 ORA를 수행하지 않았다.',
        '- log2FC/통계는 해당 비교 안에서 해석한다. baseMean은 서로 다른 처리/정규화 맥락에서 온 제공 값으로, 두 연구 간 절대 발현량 척도로 비교하지 않는다.',
        '- Cortex 통계는 submitter-provided DESeq2 statistics라고 기록되어 있고 FPKM 제출 파일에서 새로 계산한 값이 아니다. A는 count columns의 PyDESeq2 재계산으로 기록되어 있다. 원시 count 및 디자인은 이번 파일에 없어 그 분석을 다시 수행하지 않았다.',
        '', '## 재실행 및 결과 파일', '',
        '`python3 ora_analysis/analyze_deg_and_reprioritize.py`로 재현한다. 수동 선별 이유는 [우선순위 설정 TSV](../deg_aware_priority_config.tsv)에 저장했다. [별도 qPCR 요약 Markdown](../../qPCR_candidate_summary_ko.md)에는 패널과 주요 CSV 링크를 모았다.', '',
        '- [전체 Excel](DEG_aware_qPCR_and_overlap.xlsx): 전체 후보, 정확한 비교별 DEG, 완전한 심볼 overlap, 원본·중복 audit, GO 연결.',
        '- [공통 방향 52개 CSV](same_direction_52.csv), [반대 방향 23개 CSV](opposite_direction_23.csv): log2FC/gene FDR와 방향을 포함한 overlap.',
        '- [qPCR 후보 CSV](qPCR_candidates_with_DEG.csv): 이전 모든 후보와 새로운 Tmem40의 수치 및 현재 우선순위.',
        '- [피질 FDR 정렬 CSV](candidate_cortex_FDR_sort.csv): 수동 순위와 구분되는 단순 cortex FDR 정렬 참고표.',
        '- [공통 up 유의 GO 38개 CSV](common_up_GO_38.csv): 항목별 ORA 통계·공통 hit·공식 항목 링크.',
        '- [GO–gene 통계 연결 CSV](common_up_GO_with_gene_stats.csv): 공통 유의 GO 38개의 실제 공통 hit에 대한 gene-level 수치.',
        '- [해마 PBM DEG CSV](selected_hippocampus_PBM.csv), [전전두피질 PBM DEG CSV](selected_cortex_PBM.csv): 이번 분석에 사용한 정확한 비교별 입력.',
        '- [ORA–DEG 대조 CSV](ORA_DEG_alignment.csv), [전체 비교 요약 CSV](DEG_group_summary.csv): 원본 비교와 이전 ORA 입력·방향 확인.',
        '- [새 overlap CSV](new_unannotated_overlap.csv), [전체 심볼 합집합 CSV](all_selected_symbol_union.csv): ORA hit에서 빠졌던 공통 유전자와 전체 비교 범위.',
        '- [중복 심볼 CSV](duplicate_symbols_audit.csv), [미지정 심볼 CSV](placeholder_symbols_audit.csv): 원본 행을 보존한 audit.',
        '- [원본 DEG CSV](../input/DEG_FDR05_GSE245658_GSE325653.csv), [원본 ORA CSV](../input/ORA_fullcatalog_KEGG_GO.csv).',
    ]
    (out / 'DEG_aware_qPCR_report_ko.md').write_text('\n'.join(report) + '\n', encoding='utf-8')
    summary = {
        'source_sha256': hashlib.sha256(input_path.read_bytes()).hexdigest(), 'source_rows': len(raw),
        'selected_A_PBM_rows': len(aa), 'selected_A_valid_unique_symbols': len(by_gene[A]),
        'selected_A_pathway_eligible_rows': int(aa.pathway_eligible.sum()),
        'selected_B_PBM_rows': len(bb), 'selected_B_unique_symbols': len(by_gene[B]),
        'same_direction_symbols': len(same), 'same_up': int((same.A_direction == 'up').sum()),
        'same_down': int((same.A_direction == 'down').sum()), 'opposite_direction_symbols': len(opposite),
        'all_selected_symbol_union': len(joined_df), 'new_same_direction_symbols': ['Tmem40'],
        'ORA_declared_counts_and_visible_hits_match': True, 'cortex_first_panel': list(core.gene),
        'target_experiment_from_user': 'B6 cortex after tPBM; source DEG statistics are from disease/hypoxia models',
        'candidate_count': len(candidates), 'GO_shared_hit_gene_count': go_pairs.gene.nunique(),
        'GO_gene_pair_rows': len(go_pairs), 'no_new_DE_or_ORA_test': True,
    }
    (out / 'validation_summary.json').write_text(json.dumps(summary, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    print('OUTPUT', out)


if __name__ == '__main__':
    main()
