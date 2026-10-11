"""Collect validated actual atlas tables in separate MERFISH/RNA-seq sheets."""
from pathlib import Path
import json
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def main():
    merfish = json.loads((ROOT / 'results/merfish/readiness_report.json').read_text())
    if not merfish.get('actual_analysis_computed'):
        raise SystemExit('Actual MERFISH composition has not completed')
    rnaseq_path = ROOT / 'results/rnaseq/status.json'
    rnaseq = json.loads(rnaseq_path.read_text()) if rnaseq_path.exists() else {}
    smartseq_path = ROOT / 'results/rnaseq/smartseq/status.json'
    smartseq = json.loads(smartseq_path.read_text()) if smartseq_path.exists() else {}
    tables = {
        'MERFISH_major': 'results/merfish/regional_descriptive_difference_major.csv',
        'MERFISH_GABA': 'results/merfish/regional_descriptive_difference_gaba_family.csv',
        'MERFISH_neuron_subclass': 'results/merfish/regional_descriptive_difference_neuron_subclass.csv',
        'MERFISH_layer_major': 'results/merfish/composition_layer_major.csv',
        'MERFISH_layer_neuron': 'results/merfish/composition_layer_subclass_neurons.csv',
        'MERFISH_sample_coverage': 'results/merfish/sample_coverage.csv',
        'MERFISH_raw_label_mapping': 'results/merfish/raw_label_display_group_mapping.csv',
        'Core22_panel_coverage': 'results/coverage/panel22_measurement_coverage.csv',
        'PBM_core22': 'results/pbm/panel22.csv',
        'PBM_common52': 'results/pbm/same_direction_52.csv',
    }
    if rnaseq.get('status') == 'descriptive_expression_complete':
        tables.update({
            'RNAseq_gene_coverage': 'results/rnaseq/candidate_gene_coverage.csv',
            'RNAseq_cell_counts': 'results/rnaseq/selected_cell_counts.csv',
            'RNAseq_subclass_expression': 'results/rnaseq/region_subclass_gene_summary.csv',
            'RNAseq_subclass_difference': 'results/rnaseq/within_subclass_region_difference.csv',
            'RNAseq_shared_clusters': 'results/rnaseq/shared_cluster_region_difference.csv',
            'RNAseq_cluster_standardized': 'results/rnaseq/shared_cluster_standardized_difference.csv',
        })
        for sheet, relative in {
            'RNAseq_cohort_expression': 'results/rnaseq/exploratory_cohort_region_subclass_gene_summary.csv',
            'RNAseq_cohort_difference': 'results/rnaseq/exploratory_cohort_within_subclass_region_difference.csv',
            'RNAseq_cohort_cluster_weighted': 'results/rnaseq/exploratory_cohort_shared_cluster_standardized_difference.csv',
        }.items():
            if (ROOT / relative).exists():
                tables[sheet] = relative
        for sheet, relative in {
            'RNAseq_core22_sensitivity': 'results/rnaseq/panel22_P58_cohort_cluster_sensitivity.csv',
            'RNAseq_core22_screening': 'results/rnaseq/panel22_screening_coverage_and_sensitivity.csv',
        }.items():
            if (ROOT / relative).exists():
                tables[sheet] = relative
    if smartseq.get('status') == 'complete_descriptive_supplement':
        tables.update({
            'SMARTseq_matched_strata': 'results/rnaseq/smartseq/shared_strata_region_difference_min10.csv',
            'SMARTseq_same_dissection_layer': 'results/rnaseq/smartseq/same_dissection_layer_region_difference_min10.csv',
            'SMARTseq_cluster_standardized': 'results/rnaseq/smartseq/shared_cluster_standardized_difference_min10.csv',
        })
    path = ROOT / 'results/MOp_VISp_Allen_comparison.xlsx'
    table_frames = {name: pd.read_csv(ROOT / relative) for name, relative in tables.items()}
    if 'RNAseq_subclass_difference' in table_frames:
        source = table_frames['RNAseq_subclass_difference']
        panel = set(pd.read_csv(ROOT / 'results/pbm/panel22.csv')['gene'])
        table_frames['RNAseq_core22_P58'] = source.loc[
            source['age'].eq('P58') & source['gene'].isin(panel)
            & source['both_regions_observed'].eq(True)
        ].copy()
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        pd.DataFrame([
            {'dataset': 'MERFISH B6J638850', 'status': merfish['status'],
             'scope': 'Single WT male animal; sampled annotated cells; no animal P values; no density'},
            {'dataset': 'CTX-HPF 10x', 'status': rnaseq.get('status', 'pending'),
             'scope': 'B6J-background male Snap25-IRES2-Cre;Ai14 reporter; RFP-positive; sampled counts are not tissue composition; taxonomy separate from MERFISH'},
            {'dataset': 'CTX-HPF matching', 'status': 'descriptive only',
             'scope': 'P58 age-matched comparison: 1 MOp donor vs 2 VISp donors. Region and preparation date fully confounded. No animal P/FDR. P57-59 cohort sheets retain broader age differences.'},
            {'dataset': 'CTX-HPF expression units', 'status': 'full-library denominator verified',
             'scope': 'Mean log1p(CP10k); denominator is all 31,053 source genes. MOp-minus-VISp difference is not log2FC. Detection fraction is raw count > 0.'},
            {'dataset': 'Shared cluster weighting', 'status': 'exploratory',
             'scope': 'Equal weights for shared clusters; rare clusters included. Different from cell-weighted subclass mean; neither isolates a batch-independent region effect.'},
            {'dataset': 'SMART-seq v4', 'status': smartseq.get('status', 'pending'),
             'scope': 'Separate reporter/sex/age/FACS-matched strata; mean log1p(exon read CPM) with all 45,768 features as denominator. Display n>=10 per region is not a significance threshold. No identical metadata strata with 10x; not a direct replication.'},
        ]).to_excel(writer, sheet_name='Read_me', index=False)
        for name, frame in table_frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.book[name]
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
    book = load_workbook(path, read_only=True)
    for name, frame in table_frames.items():
        assert book[name].max_row == len(frame) + 1, name
        assert book[name].max_column == len(frame.columns), name
    book.close()
    print(json.dumps({'workbook': str(path.relative_to(ROOT)),
                      'sheets': len(table_frames) + 1,
                      'actual_RNAseq_included': rnaseq.get('status') == 'descriptive_expression_complete',
                      'actual_SMARTseq_included': smartseq.get('status') == 'complete_descriptive_supplement'}))


if __name__ == '__main__':
    main()
