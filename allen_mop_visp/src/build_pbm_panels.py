#!/usr/bin/env python3
"""Rebuild the PBM-to-Allen screening panels from preserved raw repository inputs.

Initial import:
  python src/build_pbm_panels.py --source-repo /workspace/tPBM_opendataset
Portable rerun, using only this analysis directory:
  python src/build_pbm_panels.py

No differential-expression or enrichment test is rerun on FDR-selected rows.
No source-repository files are changed, staged, committed, or pushed.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


A = "GSE245658"
B = "GSE325653"
A_CONTRAST = "PBM effect: 5xFAD+light vs 5xFAD"
B_CONTRAST = "PBM effect: hypobaric hypoxia + PBM vs hypobaric hypoxia"
ORA_IDS = {A: "GSE245658_hippocampus_808nm", B: "GSE325653_prefrontal_1064nm"}
PLACEHOLDERS = {"", "--", "NA", "N/A"}
SOURCE_PATHS = [
    "ora_analysis/input/DEG_FDR05_GSE245658_GSE325653.csv",
    "ora_analysis/input/ORA_fullcatalog_KEGG_GO.csv",
    "ora_analysis/deg_validation/same_direction_52.csv",
    "ora_analysis/deg_validation/common_up_GO_38.csv",
    "ora_analysis/deg_validation/common_up_GO_with_gene_stats.csv",
    "ora_analysis/deg_validation/DEG_group_summary.csv",
    "ora_analysis/deg_validation/validation_summary.json",
    "ora_analysis/deg_validation/DEG_aware_qPCR_report_ko.md",
    "ora_analysis/qpcr_candidate_config.tsv",
    "ora_analysis/deg_aware_priority_config.tsv",
    "ora_analysis/gene_function_map.tsv",
    "ora_analysis/concordant_51.csv",
    "ora_analysis/analyze_deg_and_reprioritize.py",
]
PANEL = [
    ("Entpd1", "A_adenosine", "아데노신 생성", "ATP/ADP 분해 및 AMP 생성"),
    ("Nt5e", "A_adenosine", "아데노신 생성", "CD73: AMP에서 adenosine 생성"),
    ("Slc29a1", "A_adenosine", "아데노신 수송", "ENT1: 농도 구배에 따른 양방향 nucleoside 수송"),
    ("Slc29a2", "A_adenosine", "아데노신 수송", "ENT2: nucleoside 수송"),
    ("Adk", "A_adenosine", "아데노신 대사", "adenosine kinase"),
    ("Ada", "A_adenosine", "아데노신 대사", "adenosine deaminase"),
    ("Adora1", "A_adenosine", "아데노신 반응", "A1 receptor"),
    ("Adora2a", "A_adenosine", "아데노신 반응", "A2A receptor"),
    ("Grm2", "B_primary_PBM", "기존 PBM 1차", "metabotropic glutamate receptor"),
    ("Ryr2", "B_primary_PBM", "기존 PBM 1차", "세포 내 Ca2+ 방출"),
    ("Trpc5", "B_primary_PBM", "기존 PBM 1차", "양이온 채널"),
    ("Atp2b3", "B_primary_PBM", "기존 PBM 1차", "Ca2+ 수송·항상성"),
    ("Cacna1c", "B_primary_PBM", "기존 PBM 1차", "전압 의존성 Ca2+ 채널"),
    ("Grin2a", "B_primary_PBM", "기존 PBM 1차", "NMDA receptor"),
    ("Mrps6", "B_primary_PBM", "기존 PBM 1차", "미토콘드리아 ribosomal protein"),
    ("Penk", "B_primary_PBM", "기존 PBM 1차", "enkephalin precursor"),
    ("Kcnj6", "C_electrophysiology", "전기생리 연결", "GIRK2: 막전위 조절 후보"),
    ("Shisa7", "C_electrophysiology", "전기생리 연결", "GABAergic synapse 조절 후보"),
    ("Slitrk3", "C_electrophysiology", "전기생리 연결", "억제성 시냅스 조절 후보"),
    ("Dlgap1", "C_electrophysiology", "전기생리 연결", "postsynaptic scaffold"),
    ("Bcs1l", "D_mitochondria", "미토콘드리아 확장", "호흡사슬 complex III 조립"),
    ("Ppargc1a", "D_mitochondria", "미토콘드리아 확장", "대사·미토콘드리아 조절 후보"),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def git(repo: Path, *arguments: str) -> str:
    return subprocess.check_output(["git", "-C", str(repo), *arguments], text=True).strip()


def import_sources(repo: Path, input_dir: Path) -> dict:
    """Copy explicit analysis inputs only, preserving original directory structure."""
    commit = git(repo, "rev-parse", "HEAD")
    before = git(repo, "status", "--porcelain", "--untracked-files=all")
    tracked_status = git(repo, "status", "--porcelain", "--untracked-files=no")
    files = []
    for relative in SOURCE_PATHS:
        source = repo / relative
        destination = input_dir / "repository" / relative
        source_bytes = source.read_bytes()
        committed_bytes = subprocess.check_output(["git", "-C", str(repo), "show", f"{commit}:{relative}"])
        if source_bytes != committed_bytes:
            raise ValueError(f"Source differs from recorded commit: {relative}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.exists() and destination.read_bytes() != source_bytes:
            raise ValueError(f"Preserved input already differs; refusing to overwrite: {destination}")
        if not destination.exists():
            shutil.copy2(source, destination)
        files.append({
            "source_repository_path": relative,
            "preserved_path": destination.relative_to(input_dir).as_posix(),
            "sha256": sha256(destination),
            "bytes": destination.stat().st_size,
            "matches_git_commit_content": True,
            "source_url": f"https://github.com/Sonnetly/tPBM_opendataset/blob/{commit}/{relative}",
        })
    after = git(repo, "status", "--porcelain", "--untracked-files=all")
    if before != after:
        raise RuntimeError("Source repository status changed during input preservation")
    manifest = {
        "repository": "https://github.com/Sonnetly/tPBM_opendataset",
        "git_commit": commit,
        "source_branch": git(repo, "rev-parse", "--abbrev-ref", "HEAD"),
        "imported_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_tracked_status": tracked_status,
        "source_existing_untracked_files_preserved": before.splitlines(),
        "source_repository_unchanged_during_import": True,
        "files": files,
    }
    path = input_dir / "source_provenance.json"
    if path.exists():
        previous = json.loads(path.read_text(encoding="utf-8"))
        if previous["git_commit"] != commit or previous["files"] != files:
            raise ValueError("Preserved provenance differs; refusing to overwrite")
        return previous
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-repo", type=Path, help="Import source files after checking their committed content")
    parser.add_argument("--analysis-dir", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    input_dir = args.analysis_dir / "input" / "pbm"
    out = args.analysis_dir / "results" / "pbm"
    out.mkdir(parents=True, exist_ok=True)
    if args.source_repo:
        import_sources(args.source_repo, input_dir)
    provenance = json.loads((input_dir / "source_provenance.json").read_text(encoding="utf-8"))
    checks: list[dict] = []

    def check(name: str, condition: bool, detail: object = None) -> None:
        condition = bool(condition)
        checks.append({"name": name, "passed": condition, "detail": detail})
        if not condition:
            raise ValueError(f"Validation failed: {name}; {detail}")

    for record in provenance["files"]:
        check("source_hash:" + record["source_repository_path"],
              sha256(input_dir / record["preserved_path"]) == record["sha256"])
    copied = input_dir / "repository" / "ora_analysis"
    raw = pd.read_csv(copied / "input" / "DEG_FDR05_GSE245658_GSE325653.csv",
                      keep_default_na=False, float_precision="round_trip")
    raw["input_csv_row"] = np.arange(2, len(raw) + 2)
    for column in ["log2FC", "pvalue", "FDR_padj", "baseMean"]:
        raw[column] = pd.to_numeric(raw[column], errors="raise")
        check("finite_DEG:" + column, np.isfinite(raw[column]).all())
    check("raw_DEG_rows_5651", len(raw) == 5651, len(raw))
    check("imported_FDR05_valid", raw.FDR_padj.ge(0).all() and raw.FDR_padj.lt(0.05).all())
    check("log2FC_direction_consistent", ((raw.log2FC.gt(0) & raw.direction.eq("up")) |
                                         (raw.log2FC.lt(0) & raw.direction.eq("down"))).all())
    raw["pathway_eligible"] = raw.used_in_pathway_analysis.astype(str).map({"True": True, "False": False})
    check("pathway_eligibility_parsed", raw.pathway_eligible.notna().all())
    selected = {
        "hippocampus": raw.loc[raw.dataset.eq(A) & raw.tissue.eq("hippocampus") & raw.contrast.eq(A_CONTRAST)].copy(),
        "cortex": raw.loc[raw.dataset.eq(B) & raw.tissue.eq("prefrontal cortex") & raw.contrast.eq(B_CONTRAST)].copy(),
    }
    check("correct_PBM_contrasts", [len(selected[x]) for x in selected] == [657, 998])
    check("correct_wavelengths", set(selected["hippocampus"].wavelength_nm) == {808} and
          set(selected["cortex"].wavelength_nm) == {1064})
    grouped = {label: {str(g): f for g, f in frame.loc[~frame.gene_symbol.isin(PLACEHOLDERS)].groupby("gene_symbol")}
               for label, frame in selected.items()}
    check("valid_unique_symbols", [len(grouped[x]) for x in selected] == [651, 998])
    for label, by_gene in grouped.items():
        check("duplicate_directions_consistent:" + label, all(f.direction.nunique() == 1 for f in by_gene.values()))

    def stats(gene: str) -> dict:
        result = {"gene": gene}
        directions = []
        for label in selected:
            rows = grouped[label].get(gene)
            absent = rows is None
            ambiguous = not absent and len(rows) > 1
            status = "missing_from_FDR05_table" if absent else "ambiguous_duplicate_symbol" if ambiguous else "selected_FDR05"
            direction = None if absent else str(rows.direction.iloc[0])
            directions.append(direction)
            result.update({f"{label}_status": status, f"{label}_present_in_FDR05_table": not absent,
                           f"{label}_direction": direction, f"{label}_row_count": 0 if absent else len(rows)})
            for column in ["log2FC", "FDR_padj", "baseMean", "gene_id", "input_csv_row"]:
                result[f"{label}_{column}"] = None if absent or ambiguous else rows[column].iloc[0]
            result[f"{label}_FC_treatment_over_control"] = None if absent or ambiguous else 2 ** float(rows.log2FC.iloc[0])
        result["PBM_overlap_class"] = ("same_direction" if directions[0] == directions[1] else "opposite_direction") \
            if all(directions) else "both_missing_from_selected_tables" if not any(directions) else "one_selected_table_only"
        return result

    union = pd.DataFrame([stats(g) for g in sorted(set(grouped["hippocampus"]) | set(grouped["cortex"]))])
    same = union.loc[union.PBM_overlap_class.eq("same_direction")].copy()
    opposite = union.loc[union.PBM_overlap_class.eq("opposite_direction")].copy()
    check("same_direction_overlap_52", len(same) == 52, len(same))
    check("common_up34_down18", same.hippocampus_direction.value_counts().to_dict() == {"up": 34, "down": 18})
    check("opposite_direction_overlap23", len(opposite) == 23)
    check("no_ambiguous_shared_symbols", same.hippocampus_row_count.eq(1).all() and same.cortex_row_count.eq(1).all())
    baseline_same = pd.read_csv(copied / "deg_validation" / "same_direction_52.csv", float_precision="round_trip")
    baseline = baseline_same.set_index("gene")
    check("same52_matches_committed_baseline_genes", set(same.gene) == set(baseline.index))
    for current, old in [("hippocampus_log2FC", "A_log2FC"), ("hippocampus_FDR_padj", "A_FDR_padj"),
                         ("cortex_log2FC", "B_cortex_log2FC"), ("cortex_FDR_padj", "B_cortex_FDR_padj")]:
        check("same52_matches_committed_statistics:" + current,
              np.array_equal(same.set_index("gene")[current].sort_index().to_numpy(), baseline[old].sort_index().to_numpy()))
    old51 = pd.read_csv(copied / "concordant_51.csv")
    check("full_DEG_adds_Tmem40_to_ORA51", set(same.gene) - set(old51.gene) == {"Tmem40"})

    ora = pd.read_csv(copied / "input" / "ORA_fullcatalog_KEGG_GO.csv", keep_default_na=False, float_precision="round_trip")
    ora["input_csv_row"] = np.arange(2, len(ora) + 2)
    check("raw_ORA_rows28092", len(ora) == 28092)
    check("unique_ORA_keys", not ora.duplicated(["dataset", "direction", "source", "set_id"]).any())
    ora_audit = []
    for label, dataset in [("hippocampus", A), ("cortex", B)]:
        for direction in ["up", "down"]:
            terms = ora.loc[ora.dataset.eq(ORA_IDS[dataset]) & ora.direction.eq(direction)]
            hits = {g for genes in terms.hits for g in genes.split(";") if g}
            eligible = selected[label].loc[selected[label].pathway_eligible & selected[label].direction.eq(direction)]
            declared_counts = set(terms.n_DEG_in_contrast.astype(int))
            check(f"ORA_input_alignment:{label}:{direction}", declared_counts == {len(eligible)} and hits <= set(eligible.gene_symbol))
            ora_audit.append({"dataset": dataset, "tissue": label, "direction": direction,
                              "ORA_declared_DEG_count": next(iter(declared_counts)), "eligible_DEG_rows": len(eligible),
                              "ORA_hit_symbol_count": len(hits), "all_hits_with_correct_DEG_direction": True})
    up_go = ora.loc[ora.source.eq("GO") & ora.direction.eq("up") & ora.padj_BH_full_catalog.lt(0.05)]
    indexed = {d: up_go.loc[up_go.dataset.eq(ORA_IDS[d])].set_index("set_id") for d in [A, B]}
    go_rows = []
    pairs = []
    for go_id in sorted(set(indexed[A].index) & set(indexed[B].index)):
        a, b = indexed[A].loc[go_id], indexed[B].loc[go_id]
        check("GO_label_alignment:" + go_id, a.set_name == b.set_name and a.go_aspect == b.go_aspect)
        shared_hits = sorted(set(filter(None, a.hits.split(";"))) & set(filter(None, b.hits.split(";"))))
        go_rows.append({"GO_id": go_id, "GO_name": a.set_name, "GO_aspect": a.go_aspect,
                        "hippocampus_ORA_FDR": a.padj_BH_full_catalog, "cortex_ORA_FDR": b.padj_BH_full_catalog,
                        "shared_hit_gene_count": len(shared_hits), "shared_hit_genes": ";".join(shared_hits),
                        "hippocampus_all_hits": a.hits, "cortex_all_hits": b.hits,
                        "hippocampus_ORA_input_csv_row": a.input_csv_row, "cortex_ORA_input_csv_row": b.input_csv_row,
                        "reference_url": "https://www.ebi.ac.uk/QuickGO/term/" + go_id})
        for gene in shared_hits:
            row = stats(gene)
            check(f"GO_hit_is_common_up:{go_id}:{gene}", row["PBM_overlap_class"] == "same_direction" and row["hippocampus_direction"] == "up")
            pairs.append({"GO_id": go_id, "GO_name": a.set_name, "gene": gene,
                          "hippocampus_ORA_FDR": a.padj_BH_full_catalog, "cortex_ORA_FDR": b.padj_BH_full_catalog,
                          **{k: v for k, v in row.items() if k.endswith(("log2FC", "FDR_padj"))}})
    go = pd.DataFrame(go_rows)
    pair_table = pd.DataFrame(pairs)
    go_genes = sorted(set(pair_table.gene))
    check("shared_up_GO38", len(go) == 38)
    check("GO_BP29_MF9", go.GO_aspect.value_counts().to_dict() == {"biological_process": 29, "molecular_function": 9})
    check("GO_shared_hits19_pairs131", len(go_genes) == 19 and len(pair_table) == 131)
    check("GO_one_term_without_shared_hits", go.shared_hit_gene_count.eq(0).sum() == 1)
    baseline_go = pd.read_csv(copied / "deg_validation" / "common_up_GO_38.csv", float_precision="round_trip").set_index("GO_id")
    check("GO38_matches_committed_baseline", set(go.GO_id) == set(baseline_go.index))
    for r in go.to_dict("records"):
        previous = baseline_go.loc[r["GO_id"]]
        check("GO_statistics_and_hits_match_baseline:" + r["GO_id"],
              r["hippocampus_ORA_FDR"] == previous.A_ORA_FDR and r["cortex_ORA_FDR"] == previous.B_ORA_FDR and
              r["shared_hit_genes"] == str(previous.shared_up_genes if pd.notna(previous.shared_up_genes) else ""))

    panel = pd.DataFrame([{**stats(g), "axis": axis, "axis_label_ko": label, "candidate_role_ko": role,
                           "screening_order": i + 1, "selection_basis": "user_analysis_plan",
                           "hypothesis_role_is_not_an_observed_regional_effect": True}
                          for i, (g, axis, label, role) in enumerate(PANEL)])
    check("panel22_unique", len(panel) == 22 and panel.gene.nunique() == 22)
    check("panel22_no_ambiguous_symbol_statistics", panel.hippocampus_row_count.le(1).all() and panel.cortex_row_count.le(1).all())
    check("panel_axis_sizes_8_8_4_2", panel.axis.value_counts().to_dict() == {
        "A_adenosine": 8, "B_primary_PBM": 8, "C_electrophysiology": 4, "D_mitochondria": 2})
    for gene, expected_fc, expected_fdr in [("Nt5e", -0.5367, 0.00745), ("Adora2a", -1.2929, 0.00932)]:
        row = panel.set_index("gene").loc[gene]
        check("cortex_adenosine_candidate:" + gene, row.cortex_direction == "down" and row.cortex_log2FC == expected_fc and
              np.isclose(row.cortex_FDR_padj, expected_fdr, rtol=0.005))
    missing_adenosine = panel.loc[panel.axis.eq("A_adenosine") & panel.PBM_overlap_class.eq("both_missing_from_selected_tables"), "gene"].tolist()
    check("six_adenosine_candidates_unobserved_in_FDR05", set(missing_adenosine) == {"Entpd1", "Slc29a1", "Slc29a2", "Adk", "Ada", "Adora1"})
    check("missing_values_are_not_zero_effect", panel.loc[~panel.cortex_present_in_FDR05_table, "cortex_log2FC"].isna().all())

    groups = {
        "panel22": panel.gene.tolist(),
        **{axis: panel.loc[panel.axis.eq(axis), "gene"].tolist() for axis in panel.axis.unique()},
        "same_direction52": same.gene.tolist(),
        "common_up34": same.loc[same.hippocampus_direction.eq("up"), "gene"].tolist(),
        "common_down18": same.loc[same.hippocampus_direction.eq("down"), "gene"].tolist(),
        "common_GO_hit19": go_genes,
    }
    exploratory_genes = sorted(set(panel.gene) | set(same.gene))
    exploration = pd.DataFrame([stats(g) for g in exploratory_genes])
    exploration["in_panel22"] = exploration.gene.isin(panel.gene)
    exploration["in_same_direction52"] = exploration.gene.isin(same.gene)
    exploration["in_common_GO_hit19"] = exploration.gene.isin(go_genes)
    check("combined_screening_genes62", len(exploration) == 62)
    audits = pd.concat([f.loc[f.gene_symbol.isin(PLACEHOLDERS) | f.duplicated("gene_symbol", keep=False)] for f in selected.values()], ignore_index=True)
    tables = {"panel22": panel, "same_direction_52": same, "common_up34": same.loc[same.hippocampus_direction.eq("up")],
              "common_down18": same.loc[same.hippocampus_direction.eq("down")], "opposite_direction23": opposite,
              "common_up_GO_38": go, "common_GO_hit19": pd.DataFrame([stats(g) for g in go_genes]),
              "common_GO_gene_pairs": pair_table, "exploration_genes62": exploration,
              "all_selected_symbol_union": union, "duplicate_placeholder_audit": audits,
              "ORA_DEG_alignment": pd.DataFrame(ora_audit)}
    for name, table in tables.items():
        table.to_csv(out / (name + ".csv"), index=False, encoding="utf-8-sig")
    excel_tables = {name: tables[name] for name in ["panel22", "same_direction_52", "common_up34", "common_down18",
                                                  "common_up_GO_38", "common_GO_hit19", "common_GO_gene_pairs",
                                                  "exploration_genes62", "ORA_DEG_alignment", "duplicate_placeholder_audit"]}
    readme = pd.DataFrame([
        ("source_commit", provenance["git_commit"]),
        ("scope", "PBM source verification and Atlas candidate panel preparation. Allen regional expression differences have not been calculated here."),
        ("missing_FDR05", "Missing from the FDR-selected table does not mean zero expression, zero effect, or a known nonsignificant result."),
        ("PBM_hippocampus", "5xFAD + light vs 5xFAD; 808 nm; disease model"),
        ("PBM_cortex", "Hypobaric hypoxia + PBM vs hypobaric hypoxia; prefrontal cortex; 1064 nm"),
        ("FDR", "Import original gene and GO FDR in separate columns; no BH recalculation on selected rows."),
        ("candidate_axes", "Hypothesis categories from the user plan, not observed MOp-VISp regional effects."),
    ], columns=["item", "description"])
    workbook_path = out / "PBM_panel22_and_gene_sets.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="READ_ME", index=False)
        for name, table in excel_tables.items():
            table.to_excel(writer, sheet_name=name[:31], index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 40
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for column in ws.columns:
                title = str(column[0].value)
                width = 65 if any(token in title for token in ["description", "hits", "role", "name"]) else 28
                ws.column_dimensions[column[0].column_letter].width = width
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    for name, table in excel_tables.items():
        sheet = workbook[name[:31]]
        check("excel_row_count:" + name, sheet.max_row - 1 == len(table))
        if "gene" in table.columns:
            column_number = table.columns.get_loc("gene") + 1
            exported_genes = [row[0] for row in sheet.iter_rows(min_row=2, min_col=column_number, max_col=column_number, values_only=True)]
            check("excel_gene_identity:" + name, exported_genes == table.gene.tolist())
    workbook.close()
    manifest = {"source_git_commit": provenance["git_commit"], "purpose": "Atlas baseline screening; not PBM response estimation",
                "groups": groups, "gene_counts": {name: len(genes) for name, genes in groups.items()},
                "combined_screening_gene_count": len(exploration),
                "missing_from_FDR05_table": "Not selected in provided table; full tested-gene results unavailable. Not zero expression/effect.",
                "statistics": "Supplied gene FDR and log2FC, original full-catalog ORA FDR; no new DE/ORA test or BH recalculation"}
    (out / "gene_group_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary = {"status": "passed", "source_git_commit": provenance["git_commit"], "source_hashes_verified": True,
               "raw_DEG_rows": len(raw), "raw_ORA_rows": len(ora), "selected_hippocampus_rows": len(selected["hippocampus"]),
               "selected_cortex_rows": len(selected["cortex"]), "same_direction_symbols": len(same), "same_up": 34, "same_down": 18,
               "opposite_direction_symbols": len(opposite), "shared_up_GO_terms": len(go), "common_GO_hit_genes": len(go_genes),
               "GO_gene_pairs": len(pair_table), "panel_genes": len(panel), "combined_screening_genes": len(exploration),
               "adenosine_genes_absent_from_both_selected_tables": missing_adenosine,
               "source_repository_files_modified": False, "no_new_DE_or_ORA_test": True,
               "Allen_regional_differences_calculated": False,
               "planned_scope": {"primary_Atlas_screening_panel": 22, "exploratory_same_direction_panel": 52,
                                 "inference": "Candidates only; Atlas regional effects require actual Allen measurements"},
               "observed_source_counts": {"same_direction": len(same), "common_up": 34, "common_down": 18,
                                          "GO_terms": len(go), "actual_common_GO_hits": len(go_genes)},
               "table_counts": {name: {"rows": len(table),
                                       "distinct_gene_symbols": int(table.gene.nunique()) if "gene" in table.columns else
                                           int(table.gene_symbol.nunique()) if "gene_symbol" in table.columns else None,
                                       "distinct_GO_terms": int(table.GO_id.nunique()) if "GO_id" in table.columns else None}
                                for name, table in tables.items()},
               "checks_passed": len(checks), "checks": checks,
               "output_sha256": {f.name: sha256(f) for f in sorted(out.glob("*.csv"))}}
    (out / "validation_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    adeno = panel.set_index("gene").loc[["Nt5e", "Adora2a"]]
    report = [
        "# 기존 PBM 결과와 Allen 후보 패널 연결", "",
        f"기준 저장소 commit: `{provenance['git_commit']}`. 원본 입력 13개를 `input/pbm/repository/`에 보존하고 SHA-256과 Git commit 내용을 검증했다.", "",
        "같은 방향 공통 DEG는 **52개(up 34/down 18)**, 반대 방향은 23개다. 두 up ORA에서 모두 유의한 GO는 **38개(BP29/MF9)**이고 실제 공통 hit 합집합은 **19개**다. GO 1개에는 실제 공통 hit가 없다. 전체 DEG를 사용하면 기존 ORA hit 기반 51개에 Tmem40 공통 down이 추가된다.", "",
        "22개 screening panel은 아데노신 8개, 기존 PBM 1차 8개, 전기생리 4개, 미토콘드리아 2개다. 공통 52개와 합친 탐색 범위는 62개다. 이 패널은 Atlas 기저 발현을 확인하기 위한 후보 목록이며 아직 MOp–VISp 차이가 확인된 결과나 최종 qPCR 순위는 아니다.", "",
        "| Gene | 피질 PBM log2FC | 원본 gene FDR |", "|---|---:|---:|",
        *[f"| {g} | {adeno.loc[g].cortex_log2FC:+.4f} | {adeno.loc[g].cortex_FDR_padj:.8g} |" for g in adeno.index], "",
        "Entpd1, Slc29a1, Slc29a2, Adk, Ada, Adora1은 두 비교의 제공 FDR05 표에 없다. 이는 미발현·변화 없음·유의하지 않음의 확정 결과가 아니다. `missing_from_FDR05_table`과 빈 수치로 보존하고 0으로 대체하지 않았다. 전체 tested-gene 통계가 있어야 비선택 결과를 확인할 수 있다.", "",
        "해마는 5xFAD + light vs 5xFAD(808 nm), 피질은 hypobaric hypoxia + PBM vs hypobaric hypoxia(1064 nm)이다. 정상 B6의 MOp–VISp 기저 차이 또는 지역별 PBM 효과로 직접 해석하지 않는다. Atlas 발현이 높다는 사실만으로 PBM 반응이 크거나 adenosine 순축적이 많다고 판정할 수 없다.", "",
        "원본의 gene FDR/log2FC와 GO 항목 FDR를 각각 보존했다. FDR05 subset에서 BH를 다시 계산하거나 새로운 ORA를 실행하지 않았다. 중복 심볼의 최소 FDR/최대 효과를 골라 쓰지 않고 개별 행을 audit에 보존했다. 공유 52개와 panel22에는 중복 심볼의 수치 모호성이 없다.", "",
        "재계산한 52개와 GO38의 심볼·통계·hit가 저장소 기존 결과와 일치했고, ORA 입력 개수와 모든 hit의 DEG 방향을 확인했다. 계획서의 Slitrk3 피질 FDR `0.01…`은 반올림 표기이며 결과 CSV는 원본 정밀도를 사용한다.", "",
        "재실행: `python src/build_pbm_panels.py` (pandas, numpy, openpyxl 필요). 첫 입력 보존만 `--source-repo /workspace/tPBM_opendataset`을 지정한다. 이 분석 명령은 보존된 원본을 읽고 결과 파일을 생성하며 Git 커밋·푸시는 별도 작업이다.", "",
        "- `PBM_panel22_and_gene_sets.xlsx`: 주요 패널/공통 유전자/GO/검증 결과. 저장 후 모든 sheet 행 수와 gene 순서를 다시 읽어 검증했다.",
        "- `panel22.csv`: gene별 가설 축·역할·해마/피질 방향/log2FC/FDR/status.",
        "- `same_direction_52.csv`, `common_up34.csv`, `common_down18.csv`: 탐색 공통 목록.",
        "- `common_up_GO_38.csv`, `common_GO_hit19.csv`, `common_GO_gene_pairs.csv`: GO와 실제 hit, 서로 분리한 항목/gene FDR.",
        "- `exploration_genes62.csv`, `gene_group_manifest.json`: Allen 입력용 통합 유전자와 그룹.",
        "- `duplicate_placeholder_audit.csv`, `ORA_DEG_alignment.csv`, `validation_summary.json`: 원본 audit 및 검증 기록.", "",
    ]
    (out / "PBM_bridge_report_ko.md").write_text("\n".join(report), encoding="utf-8")
    print(json.dumps({k: summary[k] for k in ["status", "checks_passed", "panel_genes", "same_direction_symbols", "shared_up_GO_terms", "common_GO_hit_genes"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
