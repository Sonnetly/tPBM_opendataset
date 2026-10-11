"""Cross-check published response snapshots and results without analysis imports."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / 'input/connectivity'
RESULTS = ROOT / 'results/connectivity'


def validate():
    checks = []

    def check(name, condition):
        if not condition:
            raise AssertionError(name)
        checks.append(name)

    def close(a, b):
        return bool(np.allclose(np.asarray(a, dtype=float), np.asarray(b, dtype=float),
                                rtol=2e-10, atol=1e-14, equal_nan=True))

    manifest = json.loads((INPUT / 'acquisition_manifest.json').read_text())
    responses = {}
    hashes = {}
    for item in manifest['base_responses'] + manifest['batch_responses']:
        path = ROOT / item['path']
        actual_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        check('raw SHA256 ' + path.name, actual_hash == item['sha256'])
        data = json.loads(path.read_text())
        check('complete response ' + path.name, data['success'] and
              len(data['msg']) == data['num_rows'] == data['total_rows'])
        check('HTTPS provenance ' + path.name, item['url'].startswith('https://api.brain-map.org/'))
        responses[path.name] = data['msg']
        hashes[item['path']] = actual_hash
    structures = {s['id']: s for s in responses['ontology_structure_graph1.json']}
    raw_experiments = responses['api_connectivity_experiments.json']
    raw_manual = {m['id']: m for name, rows in responses.items()
                  if name.startswith('manual_metadata_') for m in rows}
    raw_injection = pd.DataFrame([r for name, rows in responses.items()
                                 if name.startswith('injection_unionizes_') for r in rows])
    raw_target = pd.DataFrame([r for name, rows in responses.items()
                              if name.startswith('target_unionizes_') for r in rows])
    definitions = json.loads((INPUT / 'nrem_target_definitions.json').read_text())
    if isinstance(definitions, dict):
        definitions = definitions['targets']
    mapped = {t['acronym']: t['structure_id'] for t in definitions if t['structure_id'] is not None}
    check('exact target ontology IDs', all(structures[sid]['acronym'] == a for a, sid in mapped.items()))
    check('RT thalamic; TRN excluded; PZ unresolved', mapped.get('RT') == 262 and
          574 not in mapped.values() and any(t['acronym'] == 'PZ' and t['structure_id'] is None for t in definitions))
    ids = {}
    originals = {}
    for raw in raw_experiments:
        path = structures[raw['structure_id']]['structure_id_path'].split('/')
        region = next((a for a, sid in [('MOp', 985), ('VISp', 385)] if str(sid) in path), None)
        if raw['product_id'] == 5 and region:
            ids[raw['data_set_id']] = region
            originals[raw['data_set_id']] = raw
    metadata = pd.read_csv(RESULTS / 'selected_experiment_metadata_and_injection_audit.csv').set_index('experiment_id')
    check('selected IDs reconstructed from whole official API', set(metadata.index) == set(ids) and len(ids) == 168)
    check('source regions reconstructed from ontology', all(metadata.loc[e, 'source_region'] == a for e, a in ids.items()))
    inj = raw_injection.set_index(['section_data_set_id', 'structure_id', 'hemisphere_id'])
    check('unique raw injection keys', inj.index.is_unique)
    for eid, row in metadata.iterrows():
        raw = originals[eid]; manual = raw_manual[eid]; donor = manual['specimen']['donor']
        injections = manual['specimen']['injections']
        is_cre = bool(raw.get('transgenic_line') or raw.get('transgenic_line_id') or donor.get('transgenic_mouse_id'))
        cohort = 'Cre_line' if is_cre else ('WT_C57BL6J' if raw.get('strain') == donor.get('strain') == 'C57BL/6J' else 'non_Cre_other_or_unknown_strain')
        check(f'cohort/QC/donor {eid}', row['cohort'] == cohort and row['atlas_quality_pass'] == (manual['failed'] is False) and row['donor_id'] == donor['id'])
        source_id = 985 if ids[eid] == 'MOp' else 385
        whole = inj.loc[(eid, 997, 3), 'projection_volume']
        fraction = inj.loc[(eid, source_id, 3), 'projection_volume'] / whole
        check(f'original injection volume and purity {eid}', close(whole, raw['injection_volume']) and close(fraction, row['injection_source_parent_fraction']))
        declared = [int(s) for s in raw['injection_structures'].split('/') if s]
        source_only = bool(declared) and all(str(source_id) in structures[s]['structure_id_path'].split('/') for s in declared)
        p56 = bool(injections) and all(i.get('age', {}).get('days') == 56 for i in injections)
        check(f'source-only labels and P56 {eid}', source_only == row['source_only_annotated_injection_structures'] and p56 == row['injection_age_P56_only'])
    wt = metadata[metadata.cohort.eq('WT_C57BL6J')]
    check('WT sample counts and distinct donor IDs', wt.groupby('source_region').size().to_dict() == {'MOp': 8, 'VISp': 33} and wt.donor_id.nunique() == 41)
    targets = pd.read_csv(RESULTS / 'projection_by_experiment_target_hemisphere.csv')
    observed = targets[targets.record_observed].set_index(['experiment_id', 'target_structure_id', 'hemisphere_id']).sort_index()
    raw = raw_target.set_index(['section_data_set_id', 'structure_id', 'hemisphere_id']).sort_index()
    expected_keys = {(eid, sid, h) for eid in ids for sid in mapped.values() for h in (1, 2, 3)}
    check('all target hemisphere records present exactly once', raw.index.is_unique and observed.index.is_unique and set(raw.index) == set(observed.index) == expected_keys and len(raw) == 3528)
    metrics = ['normalized_projection_volume', 'projection_volume', 'projection_density', 'projection_energy', 'projection_intensity', 'volume', 'max_voxel_density', 'sum_pixels', 'sum_projection_pixels']
    for metric in metrics:
        check('raw metric preserved ' + metric, close(observed[metric], raw[metric]))
    denominators = np.array([originals[e]['injection_volume'] for e, _, _ in raw.index])
    check('NPV denominator is total injection signal volume', bool(np.allclose(raw.projection_volume / denominators, raw.normalized_projection_volume, rtol=2e-6, atol=1e-14)))
    pz = targets[targets.target_acronym.eq('PZ')]
    check('PZ missing anatomy preserved as unavailable metrics', len(pz) == 504 and not pz.record_observed.any() and pz[metrics].isna().all().all() and pz.measurement_status.eq('unmapped_anatomical_target').all())
    bilateral = raw_target[raw_target.hemisphere_id.eq(3)].merge(metadata.reset_index(), left_on='section_data_set_id', right_on='experiment_id')
    summary = pd.read_csv(RESULTS / 'WT_target_projection_descriptive_summary.csv')
    def subset(scope, frame):
        if scope == 'all_primary_source_injections': return frame
        if scope == 'source_only_annotated_injection_structures': return frame[frame.source_only_annotated_injection_structures]
        if scope == 'exact_P56_age': return frame[frame.injection_age_P56_only]
        if scope == 'source_parent_signal_fraction_at_least_0.8_exploratory': return frame[frame.injection_source_parent_fraction >= .8]
        if scope == 'exact_P56_and_source_parent_signal_fraction_at_least_0.8_exploratory': return frame[frame.injection_age_P56_only & (frame.injection_source_parent_fraction >= .8)]
        raise AssertionError('unknown scope ' + scope)
    for number, row in summary.iterrows():
        meta_group = subset(row.injection_scope, wt[wt.source_region.eq(row.source_region)])
        signal = bilateral[bilateral.experiment_id.isin(meta_group.index) & bilateral.structure_id.eq(mapped.get(row.target_acronym))]
        npv = signal.normalized_projection_volume
        wanted = [len(meta_group), len(signal), int((signal.projection_volume > 0).sum()), int((signal.projection_volume == 0).sum()),
                  npv.median(), npv.quantile(.25), npv.quantile(.75), npv.min(), npv.max(), signal.projection_volume.median(), signal.projection_density.median(), meta_group.injection_source_parent_fraction.median()]
        fields = ['n_experiments','n_records_observed','n_positive_segmented_signal','n_observed_zero','median_NPV','q25_NPV','q75_NPV','min_NPV','max_NPV','median_projection_volume_mm3','median_projection_density','median_injection_source_parent_fraction']
        check(f'WT summary from original target signal {number}', close(row[fields], wanted))
    differences = pd.read_csv(RESULTS / 'WT_exploratory_regional_descriptive_difference.csv')
    for number, row in differences.iterrows():
        selected = summary[summary.injection_scope.eq(row.injection_scope) & summary.target_acronym.eq(row.target_acronym)].set_index('source_region')
        m, v = selected.loc['MOp', 'median_NPV'], selected.loc['VISp', 'median_NPV']
        ratio = m / v if pd.notna(v) and v > 0 else np.nan
        check(f'exploratory median difference {number}', close([row.MOp_median_NPV, row.VISp_median_NPV, row.MOp_minus_VISp_median_NPV, row.MOp_to_VISp_median_NPV_ratio], [m, v, m-v, ratio]))
    cre_summary = pd.read_csv(RESULTS / 'Cre_line_separate_exploratory_summary.csv')
    for number, row in cre_summary.iterrows():
        selected = metadata[metadata.cohort.eq('Cre_line') & metadata.source_region.eq(row.source_region) & metadata.transgenic_line_name.eq(row.transgenic_line_name)]
        signal = bilateral[bilateral.experiment_id.isin(selected.index) & bilateral.structure_id.eq(mapped.get(row.target_acronym))]
        check(f'Cre line kept separate {number}', close([row.n_experiments, row.n_records_observed, row.median_NPV], [len(selected), len(signal), signal.normalized_projection_volume.median()]))
    wb_record = json.loads((RESULTS / 'workbook_validation.json').read_text())
    workbook = ROOT / wb_record['workbook']
    check('Excel exact SHA256', hashlib.sha256(workbook.read_bytes()).hexdigest() == wb_record['sha256'])
    wb = load_workbook(workbook, read_only=True, data_only=True)
    check('Excel sheet names and row counts', set(wb.sheetnames) == set(wb_record['sheet_rows']) and all(wb[name].max_row - 1 == count for name, count in wb_record['sheet_rows'].items()))
    wb.close()
    result = {'status': 'passed', 'validated_at_UTC': datetime.now(timezone.utc).isoformat(),
              'actual_data': True, 'validator_independent_of_analysis_imports': True,
              'checks_passed': len(checks), 'checks': checks,
              'selected_experiments': len(ids), 'WT_experiments': len(wt),
              'target_record_count': len(raw), 'unmapped_PZ_rows': len(pz),
              'WT_summary_rows_including_PZ': len(summary), 'Cre_line_summary_rows_including_PZ': len(cre_summary),
              'original_response_SHA256': hashes,
              'limits': 'Validates snapshot selection, arithmetic and reporting; no physiological or regional causal inference. Distinct tracer donors do not make heterogeneous injections a balanced region experiment.'}
    destination = ROOT / 'results/validation/connectivity_actual_data_review.json'
    destination.write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps({k: v for k, v in result.items() if k not in ('checks', 'original_response_SHA256')}, indent=2))


if __name__ == '__main__':
    validate()
