#!/usr/bin/env python3
"""Package actual projection tables with explicit interpretation and missingness."""
import hashlib
import json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "results/connectivity"
OUTPUT = RESULTS / "MOp_VISp_NREM_projection.xlsx"


def build():
    status = json.loads((RESULTS / "status.json").read_text())
    if status["status"] != "actual_projection_analysis_completed":
        raise ValueError("Only a completed actual-data analysis can be packaged")
    metadata = pd.read_csv(RESULTS / "selected_experiment_metadata_and_injection_audit.csv")
    readme = pd.DataFrame([
        ["What was measured", "Allen product5 anterograde viral tracer segmented axonal signal into atlas structures"],
        ["WT primary-source experiments", "MOp8, VISp33; C57BL/6J male; donor IDs unique41"],
        ["Cre experiments", "MOp23, VISp104; line-level exploratory tables kept separate from WT"],
        ["Signal interpretation", "Regional axonal labeling; synapses, postsynaptic cell identity, NREM promotion and PBM response unmeasured"],
        ["Normalized projection volume", "Target projection signal volume / total signal volume in injection site; dimensionless, not percentage or connection probability"],
        ["Hemispheres", "1left, 2right, 3bilateral aggregate. Do not sum 3 with 1/2. Injection-positive root record is observed only in2 in all selected experiments; absent injection-side records remain missing."],
        ["Target tiers", "VLPO/MEPO core; ZI conditional cell-type evidence; SI mixed cell-type evidence; MA/NDB regional context; RT NREM-related. PZ exact term unmapped."],
        ["Missing versus zero", "Record_observed=false or unmapped PZ has blank metrics. Observed numeric0 means no segmented signal in that record, not proof of absent axons."],
        ["Coverage", "3528 of3528 mapped-target×experiment×hemisphere records. PZ excluded from that denominator and remains unassessed."],
        ["Injection spread", "All primary-source injections plus source-only declared labels, P56, purity>=0.8 and P56+purity>=0.8 sensitivities. Purity cutoff is exploratory, not an official quality score."],
        ["Statistics", "Descriptive medians, ranges and quantiles; no p values, matched region effects or NREM-effect ranking."],
        ["Source provenance", "Raw HTTPS API responses and SHA256: input/connectivity/acquisition_manifest.json. Literature and atlas mapping: nrem_target_definitions.json."],
        ["Report", "results/connectivity/NREM_projection_report_ko.md"],
    ], columns=["item", "definition"])
    sheets = {"READ_ME": readme,
              "Target_definitions": pd.read_csv(ROOT / "input/connectivity/nrem_target_definitions.csv"),
              "WT_summary_and_sensitivity": pd.read_csv(RESULTS / "WT_target_projection_descriptive_summary.csv"),
              "WT_exploratory_differences": pd.read_csv(RESULTS / "WT_exploratory_regional_descriptive_difference.csv"),
              "WT_injection_audit": metadata[metadata.cohort.eq("WT_C57BL6J")],
              "Cre_injection_audit": metadata[metadata.cohort.eq("Cre_line")],
              "Cre_line_separate_summary": pd.read_csv(RESULTS / "Cre_line_separate_exploratory_summary.csv"),
              "Experiment_target_hemisphere": pd.read_csv(RESULTS / "projection_by_experiment_target_hemisphere.csv"),
              "Cohort_coverage_and_spread": pd.read_csv(RESULTS / "source_cohort_coverage_and_spread.csv")}
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            worksheet = writer.sheets[name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cells in worksheet.columns:
                width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in list(cells)[:40])
                worksheet.column_dimensions[cells[0].column_letter].width = min(48, max(13, width + 2))
    actual = load_workbook(OUTPUT, read_only=True, data_only=True)
    if set(actual.sheetnames) != set(sheets): raise ValueError("Workbook missing sheet")
    for name, frame in sheets.items():
        if actual[name].max_row != len(frame) + 1: raise ValueError(f"Workbook row count mismatch: {name}")
    actual.close()
    validation = {"status": "passed", "workbook": str(OUTPUT.relative_to(ROOT)), "bytes": OUTPUT.stat().st_size,
                  "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(), "sheet_rows": {k: len(v) for k,v in sheets.items()},
                  "completed_actual_analysis_required": True, "workbook_sheet_and_row_counts_verified": True}
    (RESULTS / "workbook_validation.json").write_text(json.dumps(validation, indent=2))
    print(json.dumps(validation))


if __name__ == "__main__":
    build()
